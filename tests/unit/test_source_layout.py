import pymupdf as fitz

from npa_ocr.ocr import PageResult, TableBlock, TextBlock, build_docx
from npa_ocr.coverage import review_report
from npa_ocr.source_layout import inspect_source, extract_native_blocks


def test_native_table_keeps_cell_addresses_and_outside_text():
    doc = fitz.open()
    page = doc.new_page(width=400, height=500)
    page.insert_text((40, 35), 'Heading')
    for y in (100, 150, 200):
        page.draw_line((40, y), (340, y))
    for x in (40, 190, 340):
        page.draw_line((x, 100), (x, 200))
    for x, y, text in ((50, 125, 'Name'), (200, 125, 'Value'),
                        (50, 175, 'Nitrogen'), (200, 175, '1.51%')):
        page.insert_text((x, y), text)
    source = inspect_source(page)
    blocks = extract_native_blocks(page, source, 144)
    table, = [b for b in blocks if isinstance(b, TableBlock)]
    assert (table.n_rows, table.n_cols) == (2, 2)
    assert {(c.row, c.col): c.text for c in table.cells} == {
        (0, 0): 'Name', (0, 1): 'Value', (1, 0): 'Nitrogen', (1, 1): '1.51%'}
    assert [b.text for b in blocks if not isinstance(b, TableBlock)] == ['Heading']
    assert table.bbox == (80, 200, 680, 400)
    doc.close()


def test_raster_insert_cannot_be_skipped_because_of_native_heading():
    doc = fitz.open()
    page = doc.new_page(width=400, height=500)
    page.insert_text((30, 30), 'Native heading with enough characters')
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 100, 100), False)
    pix.clear_with(200)
    page.insert_image(fitz.Rect(30, 60, 370, 470), pixmap=pix)
    source = inspect_source(page)
    assert extract_native_blocks(page, source, 200) is None
    result = review_report([PageResult(1, False, plain_text='Heading')],
                           source_inventory={1: source})
    assert result['coverage_issues'][0]['reason'] == 'raster_regions_skipped'
    assert not result['complete_transcription_verified']
    doc.close()


def test_missing_table_report_does_not_depend_on_unreadable_marker():
    report = review_report([PageResult(1, False, plain_text='Apparently good text')],
        source_inventory={1: dict(tables=[dict(bbox=[0, 0, 100, 100], rows=3, cols=2)],
                                 images=[])})
    assert report['summary']['unreadable_fragments'] == 0
    assert report['coverage_issues'][0]['reason'] == 'source_tables_without_output_tables'


def test_mixed_merge_preserves_native_heading_and_uses_raster_ocr_only_inside_image():
    from npa_ocr.source_layout import merge_mixed_blocks
    source = dict(images=[dict(bbox=[20, 100, 380, 480], area_ratio=.68)])
    heading = TextBlock(20, 'Exact native heading', 20, 380, 12)
    stale = TextBlock(120, 'Hidden text over raster', 30, 300, 12)
    noisy_heading = TextBlock(20, 'Bad OCR heading', 20, 380, 12)
    raster = TextBlock(120, 'Read from raster', 30, 300, 12)
    result = merge_mixed_blocks([heading, stale], [noisy_heading, raster], source, 72)
    assert result == [heading, raster]


def test_native_table_exports_as_table_even_without_ocr(tmp_path):
    from docx import Document
    from npa_ocr.ocr_table import TableCell
    table = TableBlock(100, (10, 100, 390, 200), [
        TableCell(0, 0, 1, 1, 'Heading'), TableCell(1, 0, 1, 1, '1.51%')], 2, 1)
    path = tmp_path / 'native.docx'
    build_docx([PageResult(1, False, blocks=[table], plain_text='Flattened fallback')], path)
    doc = Document(path)
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(1, 0).text == '1.51%'
    assert not any('Flattened fallback' in p.text for p in doc.paragraphs)


def test_vector_graphics_are_reported_without_unreadable_markers():
    report = review_report([PageResult(1, False, plain_text='Graph heading')],
        source_inventory={1: dict(images=[], tables=[], vector_paths_outside_tables=50)})
    assert report['coverage_issues'][0]['reason'] == 'vector_content_preservation_unverified'


def test_mixed_document_runs_ocr_and_preserves_native_heading(tmp_path, monkeypatch):
    import npa_ocr.ocr as ocr
    doc = fitz.open()
    page = doc.new_page(width=400, height=500)
    page.insert_text((30, 30), 'Exact native heading with enough characters')
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 100, 100), False)
    pix.clear_with(200)
    page.insert_image(fitz.Rect(30, 60, 370, 470), pixmap=pix)
    pdf = tmp_path / 'mixed.pdf'
    doc.save(pdf)
    doc.close()
    calls = []

    def recognize(*args, **kwargs):
        calls.append(True)
        return [TextBlock(20, 'Wrong OCR heading', 30, 380, 12),
                TextBlock(100, 'Raster body', 30, 300, 12)]

    monkeypatch.setattr(ocr, 'ocr_page', recognize)
    result = ocr.ocr_document(pdf, tmp_path/'mixed.docx', dpi=72, auto_rotate=False,
        predict_tables_fn=lambda path: {}, image_data_fn=lambda path: {},
        min_free_memory_mb=0)
    assert calls == [True]
    assert [b.text for b in result.pages[0].blocks] == [
        'Exact native heading with enough characters', 'Raster body']


def test_page_frame_is_not_exported_as_a_table():
    doc = fitz.open()
    page = doc.new_page(width=400, height=500)
    page.draw_rect(fitz.Rect(20, 20, 380, 480))
    page.draw_line((20, 440), (380, 440))
    page.draw_line((200, 440), (200, 480))
    page.insert_text((30, 50), 'Body text')
    page.insert_text((30, 460), 'Title block')
    assert inspect_source(page)['tables'] == []
    doc.close()


def test_tight_rows_do_not_duplicate_neighbouring_values():
    doc = fitz.open()
    page = doc.new_page(width=300, height=400)
    for y in (100, 112, 124):
        page.draw_line((20, y), (280, y))
    for x in (20, 150, 280):
        page.draw_line((x, 100), (x, 124))
    for x, y, text in ((30, 111, 'Carbon'), (160, 111, '57.8'),
                        (30, 123, 'Hydrogen'), (160, 123, '2.8')):
        page.insert_text((x, y), text, fontsize=11)
    blocks = extract_native_blocks(page, inspect_source(page), 72)
    table, = [b for b in blocks if isinstance(b, TableBlock)]
    assert {(c.row, c.col): c.text for c in table.cells} == {
        (0, 0): 'Carbon', (0, 1): '57.8', (1, 0): 'Hydrogen', (1, 1): '2.8'}
    doc.close()


def test_native_character_loss_is_reported_without_inventing_replacement():
    source = dict(images=[], tables=[], native_characters={'A': 1, '4': 1})
    report = review_report([PageResult(1, False, blocks=[TextBlock(10, 'A')])],
                           source_inventory={1: source})
    issue, = report['coverage_issues']
    assert issue['reason'] == 'native_text_character_mismatch'
    assert issue['missing_characters'] == {'4': 1}
    assert 'replacement' not in issue


def test_stitched_table_accounts_for_all_its_source_pages():
    table = TableBlock(10, (0, 10, 100, 100), [], 2, 2, source_pages=[1, 2])
    source = dict(images=[], tables=[dict(bbox=[0, 0, 100, 100], rows=2, cols=2)])
    report = review_report([PageResult(1, True, blocks=[table]), PageResult(2, True)],
                           source_inventory={1: source, 2: source})
    assert report['coverage_issues'] == []
