"""Тесты второго прочтения текста ответом DeepSeek-OCR (vlm_text.py) и формул в DOCX (docx_math.py) —
на страницах спайка npa-vlm-ocr (formulas/): слова Tesseract 5.5 и ответ DeepSeek записаны на nv-02
(fixtures/vlm_text: <key>.tess.json, <key>.ds.txt), без моделей и без Tesseract."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from npa_ocr.vlm_marks import Word
from npa_ocr.vlm_text import _accept, plain, reread_lines, tokens

FIXTURES = Path(__file__).parent / "fixtures" / "vlm_text"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _page(key: str):
    d = json.loads((FIXTURES / f"{key}.tess.json").read_text(encoding="utf-8"))
    lines: dict[tuple[int, int, int], list[Word]] = {}
    for i, t in enumerate(d["text"]):
        if t.strip():
            box = (d["left"][i], d["top"][i], d["left"][i] + d["width"][i], d["top"][i] + d["height"][i])
            lines.setdefault((d["block_num"][i], d["par_num"][i], d["line_num"][i]), []).append(
                Word(t, d["conf"][i], box))
    ordered = sorted(lines.values(), key=lambda ln: min(w.box[1] for w in ln))
    return ordered, (FIXTURES / f"{key}.ds.txt").read_text(encoding="utf-8"), d["width"][0], d["height"][0]


def _reread(key: str):
    lines, ds, width, height = _page(key)
    out, formulas = reread_lines(lines, ds, width, height)
    return lines, out, formulas


def _texts(lines) -> list[str]:
    return [" ".join(w.text for w in line) for line in lines]


def _rich(lines) -> list[str]:
    return [" ".join(w.rich or w.text for w in line) for line in lines]


# --- Разбор слов DeepSeek ---


def test_tokens_keep_markup_with_spaces_whole():
    assert tokens("Ставка (P<sub>сд атм</sub>), которая \\(P_{\\text{сд атм}}\\) - плата") == [
        "Ставка", "(P<sub>сд атм</sub>),", "которая", "\\(P_{\\text{сд атм}}\\)", "-", "плата"]


def test_plain_text_of_markup():
    assert plain("K<sub>пр</sub>") == "Kпр"
    assert plain("16<sup>3</sup>") == "16³"
    assert plain("\\(V_{\\text{ср}}\\)") == "V_ср"


def test_roman_only_over_lookalike_tesseract_word():
    """«I.» у DeepSeek на месте «г.» перед «МОСКВА» (228 стр. 1) — не римская цифра."""
    assert _accept(["II"], [Word("П", 90, (0, 0, 10, 10))])
    assert _accept(["IV"], [Word("ГУ", 90, (0, 0, 10, 10))])
    assert not _accept(["I."], [Word("г.", 90, (0, 0, 10, 10))])


def test_latin_letter_not_over_its_cyrillic_twin_or_translit():
    assert not _accept(["K"], [Word("К", 90, (0, 0, 10, 10))])  # двойник: вернее Tesseract
    assert not _accept(["i"], [Word("и", 90, (0, 0, 10, 10))])  # транслит
    assert _accept(["V"], [Word("У", 90, (0, 0, 10, 10))])
    assert _accept(["(Z)"], [Word("(", 90, (0, 0, 5, 10)), Word("7.)", 90, (6, 0, 20, 10))])


def test_indices_without_markup_become_markup():
    assert tokens("Sᵢᵏ - объем, C_t - затраты") == ["S<sub>i</sub><sup>k</sup>", "-", "объем,", "C<sub>t</sub>", "-",
                                                  "затраты"]


def test_glued_index_only_over_doubtful_word():
    """«Mфиатм» (M с индексом) — поверх неуверенного «Мам», но не «Bыбросы» поверх уверенного
    «выбросы» и не двойником буквы."""
    assert _accept(["Mфиатм"], [Word("Мам", 6, (0, 0, 10, 10))])
    assert not _accept(["Mфиатм"], [Word("Мам", 90, (0, 0, 10, 10))])
    assert not _accept(["Bыбросы"], [Word("выбросы", 30, (0, 0, 10, 10))])


def test_digit_split_only_at_digit_letter_border():
    w = lambda t: [Word(t, 90, (0, 0, 10, 10))]  # noqa: E731
    assert _accept(["3", "лет"], w("Злет"))
    assert _accept(["2013", "г."], w("2013г."))
    assert _accept(["2019", "г.", "№", "89"], w("2019г.№89"))
    assert not _accept(["1316-", "р"], w("1316-р"))  # распоряжение «1316-р» пишется слитно
    assert not _accept(["3", "лет"], w("Злот"))


# --- Страницы спайка ---


def test_roman_numerals_classes_of_hazard():
    """228 стр. 11: «ГУ и У классов опасности» → «IV и V»."""
    _, out, _ = _reread("228_p011")
    assert "производства черных металлов IV и V классов опасности, золошлаковых отходов" in _texts(out)


def test_superscripts_variables_and_formula_block():
    """710 стр. 3: «статьи 163» → 16³, «(Ким)» → K с индексом, «( 7.)» → (Z), «$»/«У» → S/V; строки
    Tesseract на месте формулы («7-5. х100%,») — формулой."""
    lines, out, formulas = _reread("710_p003")
    texts, rich = _texts(out), _rich(out)
    assert "Федерации в соответствии с пунктом 4 статьи 16³ Федерального закона" in texts
    assert any("(К<sub>нмм</sub>)," in r for r in rich)
    assert "5. Показатель сжигания (Z) определяется по формуле:" in texts
    assert any(t.startswith("S - объем сожженного") for t in texts)
    assert any(t.startswith("V - объем попутного") for t in texts)
    assert [f.latex for f in formulas] == ["Z = \\frac{S}{V} \\times 100\\%,"]
    assert "7-5. х100%," in _texts(lines) and "7-5. х100%," not in texts
    assert all(w.restored for line in out for w in line if w.rich)


def test_variables_split_by_tesseract():
    """710: «Мам», «М дгатм» → M с индексом (стр. 10); «$}», «\\:» → S, V с индексом и степенью
    (стр. 12); «С,» → C с индексом (стр. 13); «Злет» → «3 лет» (стр. 3)."""
    _, out, _ = _reread("710_p010")
    rich = _rich(out)
    assert any(r.startswith("M<sub>фиатм</sub> - выбросы i-го") for r in rich)
    assert any(r.startswith("M<sub>диатм</sub> - выбросы i-го") for r in rich)
    _, out, _ = _reread("710_p012")
    rich = _rich(out)
    assert any(r.startswith("S<sub>i</sub><sup>k</sup> - объем сожженного") for r in rich)
    assert any(r.startswith("V<sub>i</sub><sup>k</sup> - объем попутного") for r in rich)
    _, out, _ = _reread("710_p013")
    assert any(r.startswith("C<sub>t</sub> - затраты, понесенные пользователем") for r in _rich(out))
    _, out, _ = _reread("710_p003")
    assert "3 лет с момента превышения указанного показателя сжигания" in _texts(out)


def test_index_and_superscript_variables():
    """710 стр. 12: C с индексом t и степенью i, «1-м» → «i-м», «(1)» → «(t)»."""
    _, out, formulas = _reread("710_p012")
    rich = _rich(out)
    assert any(r.startswith("C<sub>t</sub><sup>i</sup> - затраты, понесенные i-м пользователем") for r in rich)
    assert "нефтяного газа в отчетном периоде (t);" in _texts(out)
    assert len(formulas) == 1


def test_inline_formula_and_footnote_marks():
    _, out, _ = _reread("248_p009")
    assert any("\\(O_2 = 15\\%\\)" in r for r in _rich(out))
    _, out, _ = _reread("1430_p010")
    texts = _texts(out)
    assert any(t.startswith("¹ При сбросе сточных вод") for t in texts)
    assert any(t.startswith("V. При очистке") for t in texts)


@pytest.mark.parametrize("key", ["156_p008", "191_p002"])
def test_pages_without_targets_are_unchanged(key: str):
    """156 стр. 8 — пункты «б)»/«в)», у DeepSeek «b)» (маркер списка — не переменная); 191 стр. 2 —
    обычный текст."""
    lines, out, formulas = _reread(key)
    assert out == lines and formulas == []


def test_no_ds_answer_keeps_lines():
    lines, _, width, height = _page("710_p003")
    assert reread_lines(lines, None, width, height) == (lines, [])


# --- DOCX ---


def test_docx_formula_and_rich_words(tmp_path: Path):
    from docx import Document

    from npa_ocr.ocr import FormulaBlock, PageResult, TextBlock, build_docx

    page = PageResult(page_number=1, ocr_applied=True, page_height=2000, blocks=[
        TextBlock(y=10, text="коэффициент (Кнмм), статьи 16³", low_words=(False, True, False, False),
                  rich_words=("", "(К<sub>нмм</sub>),", "", "")),
        FormulaBlock(y=50, latex="Z = \\frac{S}{V} \\times 100\\%,", bbox=(0, 50, 100, 80)),
        TextBlock(y=90, text="кислорода O_2 = 6%", rich_words=("", "\\(O_2 = 6\\%\\)")),
    ])
    out = tmp_path / "f.docx"
    build_docx([page], out)
    body = Document(str(out)).element.body
    paras = body.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")
    xml = "".join(p.xml for p in paras)
    assert f"{{{M}}}oMathPara" in "".join(el.tag for p in paras for el in p)
    assert 'w:vertAlign w:val="subscript"' in xml
    assert "<m:f>" in xml or "m:f>" in xml  # дробь S/V
    assert xml.count("<m:oMath>") + xml.count("<m:oMath ") >= 2  # блок и формула в строке


def test_docx_unparsable_formula_stays_as_text(tmp_path: Path):
    from docx import Document

    from npa_ocr.ocr import FormulaBlock, PageResult, build_docx

    page = PageResult(page_number=1, ocr_applied=True, page_height=2000,
                      blocks=[FormulaBlock(y=0, latex="\\frac{S}{", bbox=(0, 0, 1, 1))])
    out = tmp_path / "f.docx"
    build_docx([page], out)
    assert any("\\frac{S}{" in p.text for p in Document(str(out)).paragraphs)


def test_page_cache_keeps_formulas_and_markup(tmp_path: Path):
    from npa_ocr.ocr import FormulaBlock, PageResult, TextBlock, _load_page, _save_page

    page = PageResult(page_number=3, ocr_applied=True, page_height=2000, blocks=[
        TextBlock(y=1, text="a Kпр", rich_words=("", "K<sub>пр</sub>")),
        FormulaBlock(y=2, latex="Z = S", bbox=(1.0, 2.0, 3.0, 4.0)),
    ])
    _save_page(tmp_path / "p.json", page, {"m": 1})
    assert _load_page(tmp_path / "p.json", {"m": 1}).blocks == page.blocks
