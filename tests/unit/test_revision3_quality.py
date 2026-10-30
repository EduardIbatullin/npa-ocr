"""Regressions for lost definitions, ambiguous labels and table pagination."""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from npa_ocr.coverage import merge_note_readings, remove_table_duplicates
from npa_ocr.cell_reading import _mixed_labels, _code
from npa_ocr.ocr import TableBlock, TextBlock, PageResult, build_docx
from npa_ocr.ocr_table import TableCell
from npa_ocr.table_stitch import align_columns, merge_continuation


def test_similar_definitions_cannot_overwrite_one_another():
    vertical = "коэффициент частотной коррекции для общей вибрации в вертикальном направлении"
    horizontal = vertical.replace("вертикальном", "горизонтальном")
    result = _mixed_labels(f"Wd — {vertical}; Wd — {horizontal}", f"Wk — {vertical}; Wd — {horizontal}")
    assert result == f"Wk — {vertical};Wd — {horizontal}"
    assert _code("XIII", "хШ", "XI") == "XIII"


def test_missing_definitions_need_agreement_and_existing_neighbours():
    old = "CAS — регистрационный номер; класс опасности — по стандарту."
    complete = "CAS — регистрационный номер; ПДУ — предельно допустимый уровень; (К) — канцероген; класс опасности — по стандарту."
    recovered = merge_note_readings(old, complete, complete)
    assert recovered and "ПДУ — предельно допустимый уровень" in recovered and "(К) — канцероген" in recovered
    assert merge_note_readings(old, complete, old) is None
    assert merge_note_readings(old, "другая таблица; ПДУ — предельно допустимый уровень; конец", complete) is None


def test_table_text_is_deduplicated_only_inside_its_source_geometry():
    text = "II. Участок приемки сырья, готовой продукции"
    table = TableBlock(100, (10,100,900,400), [TableCell(0,0,1,2,text)], 1,2,column_edges=[0,.5,1])
    inside = TextBlock(110,text,20,850,20)
    outside = TextBlock(420,text,20,850,20)
    p = PageResult(1,True,blocks=[table,inside,outside],page_height=1000)
    assert remove_table_duplicates(p).blocks == [table,outside]
    assert p.blocks == [table,inside,outside]
    from dataclasses import replace
    unruled=replace(table,column_edges=[])
    assert remove_table_duplicates(replace(p,blocks=[unruled,inside])).blocks==[unruled,inside]


def test_subset_grid_is_continuation_and_scan_drift_does_not_add_column():
    a=[TableCell(0,0,1,1,"a")]
    b=[TableCell(0,0,1,1,"b")]
    assert align_columns(a,[0,.04,.27,.31,.41,.54,.80,1],b,[0,.04,.31,.80,1]) is not None
    result = align_columns(a,[0,.2,.4,.6,.8,1],b,[0,.2,.4,.6,.815,1])
    assert result and len(result[2]) == 6
    assert align_columns(a,[0,.2,.4,.6,.8,1],b,[0,.1,.33,.7,1]) is None


def test_name_split_at_page_boundary_becomes_one_merged_name():
    top=[TableCell(0,0,1,1,"2"),TableCell(0,1,1,1,"Мебель детская"),TableCell(0,2,1,1,"340 мм")]
    cont=[TableCell(0,0,7,1,""),TableCell(0,1,7,1,"дошкольная, ученическая (столы)"),TableCell(0,2,1,1,"400 мм")]
    cells,_,_=merge_continuation(top,1,cont,7)
    name=next(c for c in cells if c.row==0 and c.col==1)
    assert name.row_span==8 and name.text=="Мебель детская дошкольная, ученическая (столы)"
    assert next(c for c in cells if c.row==1 and c.col==2).text=="400 мм"


def test_export_keeps_caption_header_and_short_rows_together(tmp_path: Path):
    block=TableBlock(100,(0,100,1000,500),[
        TableCell(0,0,1,1,"№ п/п"),TableCell(0,1,1,1,"Наименование"),TableCell(0,2,1,1,"Формула"),
        TableCell(1,0,1,1,"12"),TableCell(1,1,1,1,"Вещество"),TableCell(1,2,1,1,"C₄H₈N₈O₈")],2,3,
        column_edges=[0,.04,.8,1],source_pages=[4])
    page=PageResult(1,True,blocks=[TextBlock(50,"Таблица 2.8",10,100,20),block],page_height=1000)
    path=tmp_path/'result.docx'
    build_docx([page],path)
    doc=Document(path)
    assert not any("Источник:" in p.text for p in doc.paragraphs)
    assert doc.tables[0].rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
    assert doc.tables[0].rows[1]._tr.trPr.find(qn('w:cantSplit')) is not None
    assert doc.tables[0].cell(1,2).text=="C₄H₈N₈O₈"
    assert doc.paragraphs[0].paragraph_format.keep_with_next
    assert doc.styles['Normal'].element.rPr.find(qn('w:lang')).get(qn('w:val'))=='ru-RU'
    build_docx([page],path,diagnostic=True)
    assert any("Источник:" in p.text for p in Document(path).paragraphs)


def test_upright_quantity_cannot_be_rotated_into_a_dash(monkeypatch):
    from PIL import Image, ImageDraw
    from npa_ocr.vlm_reread import _try_cell
    import npa_ocr.vlm_reread as reread

    image=Image.new('RGB',(100,100),'white')
    ImageDraw.Draw(image).rectangle((48,30,52,65),fill='black')
    monkeypatch.setattr(reread,'is_vertical',lambda image:True)
    calls=[]
    def ask(image,prompt):
        calls.append(image.size)
        return '1','stop',0
    result=_try_cell(ask,image,dict(box=[0,0,100,100],tess='',kind='empty',expected='number'))
    assert result['rot']==0 and result['used']=='1' and len(calls)==1
    header=_try_cell(ask,image,dict(box=[0,0,100,100],tess='',kind='empty',expected='number',numbering=True))
    assert header['used']=='1' and abs(header['rot'])==90


def test_faint_dash_is_visible_but_blank_cells_and_upright_digits_stay_distinct():
    import numpy as np
    from npa_ocr.stroke_geometry import faint_dash
    blank=np.full((65,140),255,dtype=np.uint8)
    assert not faint_dash(blank)
    dash=blank.copy();dash[16:18,63:77]=170
    assert faint_dash(dash)
    digit=blank.copy();digit[16:42,70:73]=100
    assert not faint_dash(digit)
    noise=dash.copy();noise[30:40,40:50]=100
    assert not faint_dash(noise)
    from PIL import Image
    fixtures=Path(__file__).parent/'fixtures/sanpin_2025_2'
    assert faint_dash(np.array(Image.open(fixtures/'faint-dash-border.png').convert('L')))
    assert not faint_dash(np.array(Image.open(fixtures/'source-blank-note.png').convert('L')))


def test_inline_requisites_require_two_identical_dates_and_numbers():
    from PIL import Image
    from npa_ocr.vlm_marks import Word
    from npa_ocr.requisites import restore_inline_requisites
    image=Image.new('RGB',(1000,1000),'white')
    lines=[[Word('от7Я 25905 „2,',20,(500,100,900,130))]]
    out=restore_inline_requisites(lines,image,lambda *a:('от 17.03.2025 № 2','stop',0))
    assert out[0][0].text=='от 17.03.2025 № 2' and out[0][0].restored
    answers=iter([('от 17.03.2025 № 2','stop',0),('от 17.03.2025 № 3','stop',0)])
    rejected=restore_inline_requisites(lines,image,lambda *a:next(answers))
    assert rejected[0][0].text==lines[0][0].text and rejected[0][0].restored
    rejected=restore_inline_requisites(lines,image,
        lambda *a:('OT 17.03.2024 No. 1','stop',0),
        secondary_ask=lambda *a:'от 14.03.2025 № 2')
    assert rejected[0][0].text==lines[0][0].text and rejected[0][0].restored


def test_inline_chapter_number_is_read_from_image_not_globally_replaced():
    from PIL import Image
    from npa_ocr.vlm_marks import Word
    from npa_ocr.text_recovery import recover_inline_codes
    image=Image.new('RGB',(1000,1000),'white')
    lines=[[Word('главы',99,(100,100,150,125)),Word('П',99,(155,100,175,125))]]
    out=recover_inline_codes(lines,image,lambda *a:('главы II','stop',0))
    assert out[0][1].text=='II' and out[0][1].restored
    assert recover_inline_codes(lines,image,lambda *a:('не читается','stop',0))==lines


def test_review_report_preserves_unreadable_output_and_does_not_claim_verification():
    from npa_ocr.coverage import review_report
    block=TableBlock(100,(0,100,500,200),[TableCell(0,0,1,1,'[не распознано]')],1,1,source_pages=[3])
    pages=[PageResult(3,True,blocks=[block,TextBlock(20,'сомнительное слово',low_words=(True,False)),
                                  TextBlock(30,'[не распознано]')])]
    result=review_report(pages)
    assert result['result_policy']=='automatic_only'
    assert result['complete_transcription_verified'] is False
    assert result['summary']==dict(flagged_fragments=3,unreadable_fragments=2)
    assert result['items'][0]['reason']=='unreadable_cell'
    assert result['items'][0]['source_pages']==[3]
    assert block.cells[0].text=='[не распознано]'
    assert all('replacement' not in item for item in result['items'])


def test_section_label_needs_whole_heading_agreement():
    from PIL import Image
    from npa_ocr.vlm_reread import _verify_section
    image=Image.new('RGB',(600,100),'white')
    d=dict(box=[0,0,600,35],expected='section',ds='I. П. Участок приемки сырья, готовой продукции')
    assert _verify_section(lambda *a:('II. Участок приемки сырья, готовой продукции','stop',0),image,d).startswith('II.')
    answers=iter([('II. Участок приемки сырья, готовой продукции','stop',0),
                  ('III. Участок приемки сырья, готовой продукции','stop',0)])
    assert _verify_section(lambda *a:next(answers),image,d) is None
    assert _verify_section(lambda *a:('I. П. Участок приемки сырья, готовой продукции','stop',0),image,d) is None


def test_real_roman_crops_fix_missing_strokes_without_dropping_existing_ones():
    import pytest
    pytest.importorskip('cv2')
    pytesseract=pytest.importorskip('pytesseract')
    from PIL import Image
    from npa_ocr.ocr import _resolve_tesseract_cmd
    from npa_ocr.vlm_tables import _Cell
    from npa_ocr.cell_reading import verify_cells
    command=_resolve_tesseract_cmd()
    if not command:
        import shutil
        command=shutil.which('tesseract')
    if not command:pytest.skip('Tesseract unavailable')
    pytesseract.pytesseract.tesseract_cmd=command
    fixtures=Path(__file__).parent/'fixtures/sanpin_2025_2'
    for file,previous,support,expected in [('roman-XIII.png','XI','хШ','XIII'),
                                           ('roman-VIIIv.png','VIIIв','УШв','VIIIв')]:
        crop=Image.open(fixtures/file).convert('RGB')
        image=Image.new('RGB',(crop.width+20,crop.height*2+30),'white')
        image.paste(crop,(10,10))
        cell=_Cell(10,10,10+crop.width,10+crop.height,text=previous,tess=support,expected='code')
        verify_cells(image,[cell])
        assert cell.text==expected


def test_broken_side_rule_is_not_a_real_baseline_ellipsis():
    import numpy as np
    from PIL import Image
    from npa_ocr.stroke_geometry import remove_broken_side_rule
    path=Path(__file__).parent/'fixtures/sanpin_2025_2/number-border-dots.png'
    original=np.array(Image.open(path).convert('L'))
    cleaned=np.array(remove_broken_side_rule(Image.open(path)).convert('L'))
    assert (original[:,:int(.85*original.shape[1])]==cleaned[:,:int(.85*original.shape[1])]).all()
    assert not (original==cleaned).all()
    baseline=np.full((40,100),255,dtype=np.uint8)
    for x in (87,91,95):baseline[31:33,x:x+2]=0
    assert (np.array(remove_broken_side_rule(Image.fromarray(baseline)).convert('L'))==baseline).all()
    from npa_ocr.response_validation import matches_cell_type
    assert matches_cell_type('15','number')
    assert not matches_cell_type('15...','number')
