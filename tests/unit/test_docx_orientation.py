from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

from npa_ocr.ocr import PageResult, TableBlock, TextBlock, build_docx, _save_page, _load_page
from npa_ocr.ocr_table import TableCell


def test_docx_follows_portrait_landscape_portrait_source_geometry(tmp_path):
    def page(number,width,height):
        table=TableBlock(100,(0,100,width,200),[
            TableCell(0,0,1,1,'Name'),TableCell(0,1,1,1,'Value'),
            TableCell(1,0,1,1,'Carbon'),TableCell(1,1,1,1,'57,8')],2,2,
            column_edges=[0,.5,1])
        return PageResult(number,False,blocks=[table],page_width=width,page_height=height)
    path=tmp_path/'mixed.docx'
    build_docx([page(1,600,840),page(2,840,600),page(3,600,840)],path)
    doc=Document(path)
    assert [s.orientation for s in doc.sections]==[
        WD_ORIENT.PORTRAIT,WD_ORIENT.LANDSCAPE,WD_ORIENT.PORTRAIT]
    assert [s.page_width>s.page_height for s in doc.sections]==[False,True,False]
    widths=[sum(c.width for c in t.columns) for t in doc.tables]
    assert widths[1]>widths[0]*1.4
    assert widths[0]==widths[2]
    assert all(t.cell(1,1).text=='57,8' for t in doc.tables)
    assert not doc.element.findall('.//'+qn('w:tblHeader'))


def test_first_landscape_page_uses_first_section(tmp_path):
    path=tmp_path/'landscape.docx'
    build_docx([PageResult(1,True,blocks=[TextBlock(30,'Heading')],
                          page_width=840,page_height=600)],path)
    doc=Document(path)
    assert len(doc.sections)==1
    assert doc.sections[0].orientation==WD_ORIENT.LANDSCAPE


def test_geometry_and_applied_rotation_survive_page_cache(tmp_path):
    path=tmp_path/'page.json'
    page=PageResult(1,True,blocks=[TextBlock(100,'Heading')],
                   page_width=840,page_height=600,raster_rotation=90)
    _save_page(path,page,{})
    restored=_load_page(path,{})
    assert (restored.page_width,restored.page_height,restored.raster_rotation)==(840,600,90)
