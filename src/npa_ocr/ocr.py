"""Распознавание PDF без текстового слоя (OCR) и сборка DOCX.

Движки распределены по задачам:

- Структура таблиц — TATR (:mod:`npa_ocr.table_structure`) или PaddleOCR
  (PPStructureV3). Уверенность PaddleOCR для поиска повреждённого текста не
  годится: мусор и обрезанный печатью текст получают те же 0.9+, что и чистый.
- Текст ячеек таблиц — RapidOCR (модель PaddleOCR eslav_PP-OCRv5 в ONNX
  Runtime): та же точность, но без падений с SIGILL на CPU без AVX2.
- Текст вне таблиц — Tesseract. На чистом скане он читает почти без ошибок, а
  его уверенность в целом отделяет текст под печатью от чистого. Не всегда:
  печать может исказить слово при высокой уверенности, поэтому формула
  должности подписанта дополнительно сверяется нечётко
  (:func:`_restore_signatory_title_words`).
- С ``table_engine="vlm"`` таблицы, удаление меток и второе прочтение текста
  выполняет DeepSeek-OCR (:mod:`npa_ocr.vlm_tables`, :mod:`npa_ocr.vlm_marks`,
  :mod:`npa_ocr.vlm_text`).

Каждый движок передаётся функцией (``predict_tables_fn``, ``image_data_fn`` и
т. п.): в работе это настоящие модели, в тестах — подмены на сохранённых
ответах без загрузки моделей.
"""

from __future__ import annotations

import difflib
import gc
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from contextlib import contextmanager
from dataclasses import asdict, dataclass, field, replace
from functools import lru_cache
from pathlib import Path
from typing import Callable, Iterable, Iterator

import pypdf
from PIL import Image

from npa_ocr import (
    ai_escalation,
    docx_math,
    mark_detection,
    table_stitch,
    table_structure,
    vlm_marks,
    vlm_reread,
    vlm_tables,
    vlm_text,
)
from npa_ocr.ocr_table import TableCell, normalize_header_text, reconstruct_table
from npa_ocr.table_structure import TableStructure

Box = tuple[float, float, float, float]  # x1, y1, x2, y2

DEFAULT_DPI = 200
# Уверенность Tesseract 0–100: у чистого текста обычно 71–97, у повреждённого печатью — 0–55;
# порог посередине разрыва.
DEFAULT_CONFIDENCE_THRESHOLD = 60
# Страница без текстового слоя даёт при извлечении текста меньше 20 знаков мусора
# (пробелы, артефакты); с этого порога считается, что текстовый слой есть.
DEFAULT_MIN_PAGE_TEXT_CHARS = 20
UNRECOGNIZABLE_MARK = "[нрзб.]"


@dataclass
class TextBlock:
    """Строка текста вне таблиц.

    Attributes:
        y: Верх строки, px.
        text: Текст строки.
        x0: Левый край, px.
        x1: Правый край, px.
        height: Высота строки, px.
        low_words: Признаки неуверенного распознавания по словам
            ``text.split()``; такие слова выделяются в DOCX.
        rich_words: Разметка по словам ``text.split(" ")``: ``<sub>``/``<sup>``,
            ``\\(LaTeX\\)``; пустая строка — слово как в ``text``.
    """

    y: float
    text: str
    x0: float = 0.0
    x1: float = 0.0
    height: float = 0.0
    low_words: tuple[bool, ...] = ()
    rich_words: tuple[str, ...] = ()


@dataclass
class FormulaBlock:
    """Формула отдельной строкой; в DOCX — формула Word по центру.

    Attributes:
        y: Верх формулы, px.
        latex: Формула LaTeX без ``\\[ \\]``.
        bbox: Рамка формулы, px.
    """

    y: float
    latex: str
    bbox: Box


@dataclass
class TableBlock:
    """Распознанная таблица страницы."""

    y: float
    bbox: Box
    cells: list[TableCell]
    n_rows: int
    n_cols: int
    dropped_columns: int = 0  # столбцы без текста, удалённые как артефакт структуры (TATR)
    column_edges: list[float] = field(default_factory=list)
    source_pages: list[int] = field(default_factory=list)


PageBlock = TextBlock | TableBlock | FormulaBlock


@dataclass
class PageResult:
    """Результат распознавания страницы."""

    page_number: int
    ocr_applied: bool  # False — страница уже имела текстовый слой, OCR не запускался
    blocks: list[PageBlock] = field(default_factory=list)
    plain_text: str | None = None  # для ocr_applied=False — текст из текстового слоя как есть
    page_height: float = 0.0  # пиксели; нужна для отбора колонтитулов (0 — не отбирать)
    # первая таблица страницы дописана в таблицу предыдущей (stitch_tables): в DOCX перед ней
    # нет разрыва страницы, иначе многостраничная таблица оставила бы пустые страницы
    continues_table: bool = False
    review_issues: list[dict] = field(default_factory=list)


@dataclass
class OcrDocumentResult:
    """Результат распознавания документа: страницы, DOCX и тайминги."""

    pages: list[PageResult]
    docx_path: Path
    timings: dict = field(default_factory=dict)
    timings_path: Path | None = None


# --- Проверка памяти ---

# Пик памяти при загрузке моделей и распознавании таблиц на машине с 16 ГБ ОЗУ может
# подвесить систему. Перед каждой страницей-сканом проверяется свободная память, и при
# нехватке распознавание останавливается с понятной ошибкой. 0 — не проверять.
DEFAULT_MIN_FREE_MEMORY_MB = 2048


class InsufficientMemoryError(RuntimeError):
    """Свободной памяти меньше заданного минимума."""


def ensure_free_memory(min_free_mb: int, where: str) -> None:
    """Проверяет, что свободной памяти не меньше ``min_free_mb``.

    Args:
        min_free_mb: Минимум свободной памяти, МБ; 0 — не проверять.
        where: Где выполняется проверка (для сообщения).

    Raises:
        InsufficientMemoryError: Памяти меньше минимума.
    """
    if min_free_mb <= 0:
        return
    import psutil

    free_mb = psutil.virtual_memory().available / (1024 * 1024)
    if free_mb < min_free_mb:
        raise InsufficientMemoryError(
            f"Свободной памяти {free_mb:.0f} МБ, нужно не меньше {min_free_mb} МБ ({where}). "
            "Распознавание остановлено, чтобы не зависла система. Закройте другие программы "
            "или запустите с меньшим --min-free-mb (0 — отключить проверку)."
        )


# --- PaddleOCR: только таблицы ---


def _extract_table_blocks(table_json: dict) -> list[TableBlock]:
    """Собирает таблицы из JSON-результата PPStructureV3.

    Args:
        table_json: Результат в формате ``save_to_json`` с ключом
            ``table_res_list``; у каждой таблицы — ``cell_box_list`` и
            ``table_ocr_pred`` (``rec_boxes``, ``rec_texts``).

    Returns:
        Таблицы страницы.
    """
    blocks: list[TableBlock] = []
    for t in table_json.get("table_res_list", []):
        cell_box_list = t.get("cell_box_list") or []
        if not cell_box_list:
            continue
        rec = t["table_ocr_pred"]
        texts = [normalize_header_text(x) for x in rec["rec_texts"]]
        cells = reconstruct_table(cell_box_list, rec["rec_boxes"], texts)
        n_rows = max(c.row + c.row_span for c in cells)
        n_cols = max(c.col + c.col_span for c in cells)
        x1 = min(b[0] for b in cell_box_list)
        y1 = min(b[1] for b in cell_box_list)
        x2 = max(b[2] for b in cell_box_list)
        y2 = max(b[3] for b in cell_box_list)
        blocks.append(
            TableBlock(y=y1, bbox=(x1, y1, x2, y2), cells=cells, n_rows=n_rows, n_cols=n_cols)
        )
    return blocks


def run_paddle_tables(image_path: Path, predict_fn: Callable[[Path], dict]) -> list[TableBlock]:
    """Распознаёт таблицы страницы через PaddleOCR.

    Args:
        image_path: Изображение страницы.
        predict_fn: Распознавание, возвращающее JSON-совместимый словарь (а не
            объект PPStructureV3), чтобы код не зависел от API модели.

    Returns:
        Таблицы страницы.
    """
    return _extract_table_blocks(predict_fn(image_path))



# --- Tesseract: весь текст вне таблиц ---


def _point_in_any_box(cx: float, cy: float, boxes: list[Box]) -> bool:
    """Проверяет, лежит ли точка в одной из рамок."""
    return any(x1 <= cx <= x2 and y1 <= cy <= y2 for x1, y1, x2, y2 in boxes)


def expected_signatory_title(authority_name: str) -> str | None:
    """Формула должности подписанта для органа, издавшего акт.

    Это правило оформления документов (какая должность стоит перед названием
    органа), а не сведения о конкретном документе, поэтому восстановление по
    нему не выдумывает реквизитов. ФИО подписанта не восстанавливается.

    Args:
        authority_name: Наименование органа в именительном падеже.

    Returns:
        Формула должности или None, если для такого органа правила нет.
    """
    name = authority_name.strip()
    if name == "Президент Российской Федерации":
        return name
    if name.startswith("Правительство"):
        # родительный падеж: «Правительство» → «Правительства»
        return "Председатель Правительства" + name[len("Правительство") :]
    if name.startswith(("Федеральное агентство", "Федеральная служба")):
        return "Руководитель"
    # «Министерство X» → «Министр X» требует родительного падежа названия ведомства,
    # который надёжно не выводится из именительного; не поддерживается.
    return None


_NEAR_MISS_SIMILARITY_THRESHOLD = 0.7


_MIN_PREFIX_MATCH_LENGTH = 4


def _is_near_miss(word: str, expected_word: str) -> bool:
    """Проверяет, что слово — искажённое или обрезанное ожидаемое слово.

    Уверенность Tesseract не всегда выдаёт повреждение печатью: «Президент»
    может прочитаться как «Ярезидент» с высокой уверенностью. Сравнение целых
    слов (похожесть ≥ 0.7) ловит замену одной-двух букв. Обрезанное печатью
    слово («Прави» вместо «Правительства») сравнивается с началом ожидаемого
    слова той же длины; для слов короче четырёх букв это не делается.

    Args:
        word: Распознанное слово.
        expected_word: Ожидаемое слово.
    """
    word_l, expected_l = word.lower(), expected_word.lower()
    full_ratio = difflib.SequenceMatcher(None, word_l, expected_l).ratio()
    if full_ratio >= _NEAR_MISS_SIMILARITY_THRESHOLD:
        return True
    if len(word_l) >= _MIN_PREFIX_MATCH_LENGTH:
        prefix_ratio = difflib.SequenceMatcher(None, word_l, expected_l[: len(word_l)]).ratio()
        if prefix_ratio >= _NEAR_MISS_SIMILARITY_THRESHOLD:
            return True
    return False


def _restore_signatory_title_words(words: list[str], authority_name: str | None) -> list[str]:
    """Восстанавливает формулу должности подписанта, повреждённую печатью.

    Поиск идёт по сквозной последовательности слов страницы, а не по строкам:
    печать сдвигает геометрию слов, и Tesseract разрывает формулу на разные
    строки.

    Окно длиной с формулу заменяется, если каждое его слово — ``[нрзб.]``,
    ожидаемое слово или похожее на него (:func:`_is_near_miss`), и хотя бы одно
    совпадает точно: «Российской Федерации» встречается и в обычном тексте.
    Остальные слова, включая ФИО после формулы, не меняются.

    Args:
        words: Слова страницы вне таблиц в порядке чтения.
        authority_name: Наименование органа или None.

    Returns:
        Слова с восстановленной формулой или исходные слова.
    """
    if authority_name is None:
        return words
    expected = expected_signatory_title(authority_name)
    if expected is None:
        return words
    expected_words = expected.split()
    n = len(expected_words)
    for start in range(len(words) - n + 1):
        window = words[start : start + n]
        has_exact_match = False
        accepted = True
        for word, expected_word in zip(window, expected_words):
            if word == UNRECOGNIZABLE_MARK:
                continue
            if word.lower() == expected_word.lower():
                has_exact_match = True
                continue
            if not _is_near_miss(word, expected_word):
                accepted = False
                break
        if not accepted or not has_exact_match:
            continue
        return [*words[:start], *expected_words, *words[start + n :]]
    return words


def run_tesseract_text(
    image_path: Path,
    table_bboxes: list[Box],
    *,
    image_data_fn: Callable[[Path], dict],
    confidence_threshold: int = DEFAULT_CONFIDENCE_THRESHOLD,
    authority_name: str | None = None,
    detect_marks_fn: Callable[[Path, list[Box]], list[mark_detection.Mark]] | None = None,
    recover_covered_text_fn: Callable[[Path], str] | None = None,
    clean_marks_fn: Callable[[list[list[vlm_marks.Word]]], list[list[vlm_marks.Word]]] | None = None,
    reread_text_fn: Callable[[list[list[vlm_marks.Word]]],
                             tuple[list[list[vlm_marks.Word]], list[vlm_text.Formula]]] | None = None,
) -> list[TextBlock | FormulaBlock]:
    """Распознаёт текст страницы вне таблиц через Tesseract.

    Слова внутри таблиц пропускаются: таблицы распознаются отдельно. Слова с
    уверенностью ниже порога помечаются как неуверенные и выделяются в DOCX.

    Args:
        image_path: Изображение страницы.
        table_bboxes: Рамки таблиц страницы.
        image_data_fn: Распознавание в формате
            ``pytesseract.image_to_data(..., output_type=Output.DICT)``.
        confidence_threshold: Порог уверенности Tesseract (0–100).
        authority_name: Наименование органа; если задано, восстанавливается
            формула должности подписанта (:func:`_restore_signatory_title_words`).
        detect_marks_fn: Поиск и классификация меток
            (:func:`npa_ocr.mark_detection.detect_marks`). Текст меток
            регистрации и электронной подписи, не перекрывающих текст,
            удаляется.
        recover_covered_text_fn: Восстановление текста под меткой
            (:func:`npa_ocr.ai_escalation.recover_covered_text`).
            Восстановленные слова остаются помеченными как неуверенные.
        clean_marks_fn: Удаление текста меток по трём прочтениям
            (:func:`npa_ocr.vlm_marks.clean_lines`); используется вместо
            ``detect_marks_fn`` при ``table_engine="vlm"``.
        reread_text_fn: Второе прочтение по DeepSeek-OCR
            (:func:`npa_ocr.vlm_text.reread_lines`), выполняется после
            ``clean_marks_fn``.

    Returns:
        Строки текста и формулы-блоки сверху вниз.
    """
    data = image_data_fn(image_path)
    lines: dict[tuple[int, int, int], list[tuple[float, float, str, float, float, bool]]] = {}
    word_conf: dict[tuple[tuple[int, int, int], float, float], float] = {}  # для словесных правок
    for i in range(len(data["text"])):
        word = data["text"][i].strip()
        if not word:
            continue
        # те же исправления известных искажений OCR, что и в таблицах: смесь латиницы и
        # кириллицы, «Γ.» вместо «г.» встречаются и в обычном тексте
        word = normalize_header_text(word)
        x, y, w, h = data["left"][i], data["top"][i], data["width"][i], data["height"][i]
        cx, cy = x + w / 2, y + h / 2
        if _point_in_any_box(cx, cy, table_bboxes):
            continue
        conf = data["conf"][i]
        low = conf < confidence_threshold
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        lines.setdefault(key, []).append((y, x, word, w, h, low))
        word_conf[(key, x, y)] = float(conf)

    line_items = []  # (y, key, [(x, word), ...]) — порядок строк на странице
    for key, words in lines.items():
        words.sort(key=lambda w: w[1])  # по x внутри строки
        y = min(w[0] for w in words)
        line_items.append((y, key, words))
    line_items.sort(key=lambda item: item[0])

    rich_lines: list[tuple[str, ...]] | None = None
    formulas: list[vlm_text.Formula] = []
    if clean_marks_fn is not None or reread_text_fn is not None:
        line_items, rich_lines, formulas = _apply_word_fns(line_items, word_conf, clean_marks_fn, reread_text_fn)

    if detect_marks_fn is not None:
        line_items = _apply_mark_escalation(
            line_items,
            image_path,
            detect_marks_fn=detect_marks_fn,
            recover_covered_text_fn=recover_covered_text_fn,
        )

    flat_words = [w[2] for _, _, words in line_items for w in words]
    flat_low = [w[5] for _, _, words in line_items for w in words]
    restored_words = _restore_signatory_title_words(flat_words, authority_name)
    # слово, заменённое восстановленной формулой, уже не «сомнительное»
    flat_low = [
        low and new == old for low, old, new in zip(flat_low, flat_words, restored_words)
    ]

    blocks: list[TextBlock | FormulaBlock] = []
    pos = 0
    for i, (y, _key, words) in enumerate(line_items):
        restored = restored_words[pos : pos + len(words)]
        low_flags = tuple(flat_low[pos : pos + len(words)])
        pos += len(words)
        rich = rich_lines[i] if rich_lines is not None and any(rich_lines[i]) else ()
        blocks.append(
            TextBlock(
                y=y,
                text=" ".join(restored),
                x0=min(w[1] for w in words),
                x1=max(w[1] + w[3] for w in words),
                height=max(w[4] for w in words),
                low_words=low_flags,
                rich_words=rich,
            )
        )
    blocks.extend(FormulaBlock(y=f.box[1], latex=f.latex, bbox=f.box) for f in formulas)
    blocks.sort(key=lambda b: b.y)
    return blocks


def _apply_word_fns(
    line_items: list[tuple[float, tuple[int, int, int], list["_WordTuple"]]],
    word_conf: dict[tuple[tuple[int, int, int], float, float], float],
    clean_marks_fn: Callable[[list[list[vlm_marks.Word]]], list[list[vlm_marks.Word]]] | None,
    reread_text_fn: Callable[[list[list[vlm_marks.Word]]],
                             tuple[list[list[vlm_marks.Word]], list[vlm_text.Formula]]] | None,
) -> tuple[list[tuple[float, tuple[int, int, int], list["_WordTuple"]]], list[tuple[str, ...]],
           list[vlm_text.Formula]]:
    """Применяет к строкам Tesseract удаление меток и второе прочтение.

    Строки переводятся в слова :class:`npa_ocr.vlm_marks.Word`, обрабатываются
    и возвращаются обратно. Порядок строк сохраняется; заменённые и дописанные
    слова помечаются как неуверенные.

    Returns:
        Строки, разметка слов по строкам и формулы-блоки.
    """
    word_lines = [
        [
            vlm_marks.Word(word, word_conf.get((key, x, y), 0.0), (x, y, x + w, y + h))
            for (y, x, word, w, h, _low) in words
        ]
        for _y, key, words in line_items
    ]
    lows = {(wd.box, wd.text): low for (_y, _key, words), wl in zip(line_items, word_lines)
            for wd, (*_, low) in zip(wl, words)}
    if clean_marks_fn is not None:
        word_lines = clean_marks_fn(word_lines)
    formulas: list[vlm_text.Formula] = []
    if reread_text_fn is not None:
        word_lines, formulas = reread_text_fn(word_lines)
    result = []
    rich = []
    for i, wl in enumerate(word_lines):
        words = [
            (wd.box[1], wd.box[0], wd.text, wd.box[2] - wd.box[0], wd.box[3] - wd.box[1],
             wd.restored or lows.get((wd.box, wd.text), False))
            for wd in wl
        ]
        result.append((min(w[0] for w in words), (-1, i, 0), words))
        rich.append(tuple(wd.rich for wd in wl))
    return result, rich, formulas


_ESCALATION_CROP_PADDING = 25  # px при 200 DPI: запас вокруг метки при вырезании для ИИ-эскалации

_WordTuple = tuple[float, float, str, float, float, bool]  # y, x, word, w, h, low


def _word_center(word: _WordTuple) -> tuple[float, float]:
    """Центр слова."""
    y, x, _word, w, h, _low = word
    return x + w / 2, y + h / 2


def _apply_mark_escalation(
    line_items: list[tuple[float, tuple[int, int, int], list[_WordTuple]]],
    image_path: Path,
    *,
    detect_marks_fn: Callable[[Path, list[Box]], list[mark_detection.Mark]],
    recover_covered_text_fn: Callable[[Path], str] | None,
) -> list[tuple[float, tuple[int, int, int], list[_WordTuple]]]:
    """Применяет ИИ-эскалацию к строкам страницы (см. :func:`run_tesseract_text`).

    Метки ищутся один раз на страницу: один штамп обычно задевает несколько
    строк. Текст под меткой восстанавливается одним запросом на всю её
    область: по мелким фрагментам без контекста модель отвечает мусором.
    Полученные слова по порядку чтения раздаются неуверенным словам внутри
    метки (см. :func:`_consume_escalated_words`).

    Если модель вернула меньше слов, чем неуверенных слов в области, ответ
    считается неудачным и для этой метки не применяется: область, найденная
    по неуверенным словам, может захватить край настоящей строки, и раздать
    короткий ответ на несколько слов подряд опаснее, чем оставить их
    помеченными.
    """
    low_word_boxes = [
        (x, y, x + w, y + h) for _, _, words in line_items for (y, x, _word, w, h, low) in words if low
    ]
    print(f"[ИИ-эскалация] низкоуверенных слов на странице: {len(low_word_boxes)}", file=sys.stderr)
    if not low_word_boxes:
        return line_items
    marks = detect_marks_fn(image_path, low_word_boxes)
    print(
        f"[ИИ-эскалация] найдено областей: {len(marks)} — "
        + "; ".join(f"{m.mark_type} (covers_text={m.covers_text})" for m in marks),
        file=sys.stderr,
    )
    if not marks:
        return line_items

    discard_bboxes = [
        m.bbox
        for m in marks
        if m.mark_type in ("электронная_подпись", "штамп_регистрации") and not m.covers_text
    ]
    escalate_marks = [m for m in marks if m.covers_text]

    mark_words: list[list[str] | None] = [None] * len(escalate_marks)
    if escalate_marks and recover_covered_text_fn is not None:
        page_image = Image.open(image_path)
        for i, mark in enumerate(escalate_marks):
            words = _recover_mark_region(page_image, mark.bbox, recover_covered_text_fn)
            expected = _count_points_in_box(
                (_box_center(b) for b in low_word_boxes), mark.bbox
            )
            if words is not None and len(words) < expected:
                print(
                    f"[ИИ-эскалация] для области {mark.bbox} восстановлено "
                    f"{len(words)} слов, а низкоуверенных слов в области {expected} — "
                    "похоже на неудачное распознавание, замена не применяется",
                    file=sys.stderr,
                )
                words = None
            mark_words[i] = words
    mark_iters = [iter(words) if words else None for words in mark_words]

    result = []
    for _y, key, words in line_items:
        kept = words
        if discard_bboxes:
            kept = [w for w in kept if not _point_in_any_box(*_word_center(w), discard_bboxes)]
        if escalate_marks:
            kept = _consume_escalated_words(kept, escalate_marks, mark_iters)
        if kept:
            result.append((min(w[0] for w in kept), key, kept))
    return result


def _box_center(box: Box) -> tuple[float, float]:
    """Центр рамки."""
    return (box[0] + box[2]) / 2, (box[1] + box[3]) / 2


def _count_points_in_box(points: Iterable[tuple[float, float]], box: Box) -> int:
    """Число точек внутри рамки."""
    return sum(1 for cx, cy in points if box[0] <= cx <= box[2] and box[1] <= cy <= box[3])


def _consume_escalated_words(
    words: list[_WordTuple],
    escalate_marks: list[mark_detection.Mark],
    mark_iters: list[Iterator[str] | None],
) -> list[_WordTuple]:
    """Заменяет неуверенные слова в метках словами восстановленного текста.

    Неуверенное слово внутри метки, перекрывающей текст, заменяется следующим
    словом восстановленного для этой метки текста. Если слова кончились или
    восстановление не удалось, исходное слово остаётся. Заменённые слова
    остаются помеченными как неуверенные.
    """
    result: list[_WordTuple] = []
    for word in words:
        if not word[5]:
            result.append(word)
            continue
        cx, cy = _word_center(word)
        mark_index = next(
            (i for i, m in enumerate(escalate_marks) if m.bbox[0] <= cx <= m.bbox[2] and m.bbox[1] <= cy <= m.bbox[3]),
            None,
        )
        iterator = mark_iters[mark_index] if mark_index is not None else None
        recovered_word = next(iterator, None) if iterator is not None else None
        if recovered_word is None:
            result.append(word)
            continue
        y, x, _text, w, h, _low = word
        result.append((y, x, recovered_word, w, h, True))
    return result


def _recover_mark_region(
    page_image: Image.Image,
    mark_bbox: Box,
    recover_fn: Callable[[Path], str],
) -> list[str] | None:
    """Восстанавливает текст под меткой одним запросом на всю её область.

    Любая ошибка ``recover_fn`` подавляется: эскалация — необязательное
    улучшение и не должна ронять распознавание.

    Args:
        page_image: Изображение страницы.
        mark_bbox: Рамка метки; вырезается с запасом ``_ESCALATION_CROP_PADDING``.
        recover_fn: Восстановление текста по фрагменту.

    Returns:
        Слова восстановленного текста или None, если результата нет.
    """
    x0 = mark_bbox[0] - _ESCALATION_CROP_PADDING
    y0 = mark_bbox[1] - _ESCALATION_CROP_PADDING
    x1 = mark_bbox[2] + _ESCALATION_CROP_PADDING
    y1 = mark_bbox[3] + _ESCALATION_CROP_PADDING
    box = (
        max(0, int(x0)),
        max(0, int(y0)),
        min(page_image.width, int(x1)),
        min(page_image.height, int(y1)),
    )
    if box[2] <= box[0] or box[3] <= box[1]:
        return None
    crop = page_image.crop(box)
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
        crop.save(tmp.name)
        tmp_path = Path(tmp.name)
    try:
        text = recover_fn(tmp_path)
    except Exception as exc:
        print(f"[ИИ-эскалация] восстановление не удалось: {exc}", file=sys.stderr)
        return None
    finally:
        tmp_path.unlink(missing_ok=True)
    words = text.split()
    print(f"[ИИ-эскалация] восстановлено для области {mark_bbox}: {' '.join(words)!r}", file=sys.stderr)
    return words or None


_TABLE_CROP_PADDING = 20  # px при 200 DPI: запас вокруг таблицы при вырезании


# Область таблицы от TATR бывает обрезана справа: последний столбец не попадает во
# фрагмент, а таблица всё равно проходит is_plausible_table. Доля ширины страницы справа
# от области у верных таблиц (в том числе широких без рамок) не больше 0.17, у обрезанных —
# от 0.23; при большем отступе таблица перепроверяется через img2table.
_IMG2TABLE_RIGHT_MARGIN_THRESHOLD = 0.20

# img2table перепроверяет только принятые, но подозрительно узкие таблицы TATR. На
# отклонённых областях он сам находит ложные таблицы в обычном тексте и пользы не даёт.
# Доля перекрытия по вертикали, при которой таблица img2table считается той же, что
# область TATR (см. _best_img2table_match):
_IMG2TABLE_MIN_Y_OVERLAP_RATIO = 0.3


def _table_cells_from_img2table(table) -> tuple[list[TableCell], int, int]:
    """Переводит ``ExtractedTable`` из img2table в ячейки.

    ``table.content`` — плотная сетка, где объединённая ячейка повторяется с
    той же рамкой в каждой покрытой позиции (отдельных rowspan/colspan в
    img2table нет). Повтор рамки расширяет объединение, а не создаёт дубль.

    Returns:
        Ячейки, число строк и число столбцов.
    """
    rows_keys = sorted(table.content.keys())
    n_rows = len(rows_keys)
    n_cols = max((len(table.content[r]) for r in rows_keys), default=0)
    seen: dict[tuple, TableCell] = {}
    cells: list[TableCell] = []
    for r_idx, r_key in enumerate(rows_keys):
        row_cells = table.content[r_key]
        for c_idx, cell in enumerate(row_cells):
            if c_idx >= n_cols:
                continue
            bbox_key = (cell.bbox.x1, cell.bbox.y1, cell.bbox.x2, cell.bbox.y2)
            owner = seen.get(bbox_key)
            if owner is not None:
                owner.row_span = max(owner.row_span, r_idx - owner.row + 1)
                owner.col_span = max(owner.col_span, c_idx - owner.col + 1)
                continue
            new_cell = TableCell(row=r_idx, col=c_idx, row_span=1, col_span=1, text=(cell.value or "").strip())
            seen[bbox_key] = new_cell
            cells.append(new_cell)
    return _merge_img2table_fragment_rows(cells, n_rows, n_cols)


_FRAGMENT_CONTINUATION_START = "([«—–-,;"  # см. _looks_like_fragment_continuation


def _looks_like_fragment_continuation(text: str) -> bool:
    """Проверяет, что текст продолжает предыдущую строку ячейки.

    Название (вещества, процесса) начинается с заглавной буквы, а перенесённое
    продолжение или уточнение — со строчной буквы, скобки, кавычки, тире или
    запятой: «Углерода оксид» → «(углерод окись;» → «угарный газ)» →
    «Азота диоксид» (новое название).
    """
    stripped = text.lstrip()
    if not stripped:
        return False
    first = stripped[0]
    if first in _FRAGMENT_CONTINUATION_START:
        return True
    return first.islower()


def _merge_img2table_fragment_rows(
    cells: list[TableCell], n_rows: int, n_cols: int
) -> tuple[list[TableCell], int, int]:
    """Склеивает строки, на которые img2table раздробил многострочные ячейки.

    img2table режет таблицу по строкам текста, и перенос внутри ячейки
    становится новой строкой таблицы; по высоте и зазору такую строку не
    отличить от настоящей. Столбцы дробятся независимо, поэтому склейка идёт по
    каждому столбцу отдельно: текст, похожий на продолжение
    (:func:`_looks_like_fragment_continuation`), дописывается к предыдущему
    непустому тексту столбца. Строка удаляется, если после склейки в ней не
    осталось текста.

    Returns:
        Ячейки, число строк и число столбцов.
    """
    if n_rows <= 1:
        return cells, n_rows, n_cols

    by_row: dict[int, dict[int, TableCell]] = {}
    for cell in cells:
        by_row.setdefault(cell.row, {})[cell.col] = cell

    for c in range(n_cols):
        anchor: TableCell | None = None
        for r in range(n_rows):
            cell = by_row.get(r, {}).get(c)
            if cell is None or not cell.text.strip():
                continue
            if anchor is not None and _looks_like_fragment_continuation(cell.text):
                anchor.text = anchor.text + "\n" + cell.text
                cell.text = ""
            else:
                anchor = cell

    absorbed_rows = {
        r for r in range(n_rows) if not any(cell.text.strip() for cell in by_row.get(r, {}).values())
    }
    if not absorbed_rows:
        return cells, n_rows, n_cols

    remap: dict[int, int] = {}
    new_idx = 0
    for old_row in range(n_rows):
        if old_row in absorbed_rows:
            continue
        remap[old_row] = new_idx
        new_idx += 1

    remaining = [cell for cell in cells if cell.row not in absorbed_rows]
    for cell in remaining:
        old_end = cell.row + cell.row_span
        cell.row_span = max(1, sum(1 for rr in range(cell.row, old_end) if rr not in absorbed_rows))
        cell.row = remap[cell.row]
    return remaining, new_idx, n_cols


def _img2table_tables_for_page(image_path: Path) -> list[tuple[Box, list[TableCell], int, int]]:
    """Ищет таблицы на всей странице через img2table (линии OpenCV и Tesseract).

    Используется для перепроверки области TATR и поэтому ищет границы таблицы
    заново на всей странице: столбец, не попавший во фрагмент TATR, другим OCR
    не восстановить. Работает при 200 DPI (при 300 DPI img2table линий не
    находит). Ошибки (нет библиотеки из extra ``models``, сбой распознавания)
    не фатальны: возвращается пустой список.

    Returns:
        Правдоподобные таблицы: ``(рамка, ячейки, число строк, число столбцов)``.
    """
    try:
        from img2table.document import Image as I2TImage
        from img2table.ocr import TesseractOCR as I2TTesseractOCR
    except ImportError:
        return []
    try:
        extracted = I2TImage(src=str(image_path)).extract_tables(
            ocr=I2TTesseractOCR(n_threads=1, lang="rus"),
            implicit_rows=True,
            implicit_columns=True,
            borderless_tables=True,
            min_confidence=50,
        )
    except Exception as exc:
        print(f"[img2table] сбой fallback-распознавания {image_path}: {exc}", file=sys.stderr)
        return []
    result: list[tuple[Box, list[TableCell], int, int]] = []
    for t in extracted:
        cells, n_rows, n_cols = _table_cells_from_img2table(t)
        if cells and table_structure.is_plausible_table(cells, n_rows, n_cols):
            result.append(((float(t.bbox.x1), float(t.bbox.y1), float(t.bbox.x2), float(t.bbox.y2)), cells, n_rows, n_cols))
    return result


def _best_img2table_match(
    by1: float, by2: float, candidates: list[tuple[Box, list[TableCell], int, int]]
) -> tuple[Box, list[TableCell], int, int] | None:
    """Выбирает таблицу img2table, соответствующую области TATR по вертикали.

    Мера — перекрытие по вертикали, отнесённое к меньшей из высот; ниже
    ``_IMG2TABLE_MIN_Y_OVERLAP_RATIO`` совпадения нет.

    Returns:
        Подходящая таблица или None.
    """
    best = None
    best_ratio = 0.0
    for candidate in candidates:
        cy1, cy2 = candidate[0][1], candidate[0][3]
        overlap = min(by2, cy2) - max(by1, cy1)
        span = min(by2 - by1, cy2 - cy1)
        ratio = overlap / span if span > 0 else 0.0
        if ratio > best_ratio:
            best_ratio = ratio
            best = candidate
    return best if best_ratio >= _IMG2TABLE_MIN_Y_OVERLAP_RATIO else None


LOW_CONFIDENCE_THRESHOLD = 0.8  # уверенность строки распознавания (0–1): ниже — выделяется в DOCX


@lru_cache(maxsize=1)
def _paddle_runtime_usable() -> bool:
    """Проверяет, можно ли загрузить PaddlePaddle в этом окружении.

    CPU-сборки paddlepaddle требуют AVX2: без него процесс падает с SIGILL уже
    на ``import paddle``, и try/except это не ловит. Поэтому импорт проверяется
    в подпроцессе. Нужен только для PPStructureV3; текст ячеек читает RapidOCR.
    """
    try:
        cpuinfo = Path("/proc/cpuinfo").read_text(encoding="utf-8", errors="replace")
        if "avx2" not in cpuinfo.lower():
            return False
    except OSError:
        pass
    try:
        result = subprocess.run(
            [sys.executable, "-c", "import paddle"],
            capture_output=True,
            timeout=120,
        )
    except (OSError, subprocess.TimeoutExpired):
        return False
    return result.returncode == 0


def _default_cell_text_fn() -> Callable[[Path], tuple[str, tuple[bool, ...]]]:
    """Создаёт распознаватель фрагмента ячейки на RapidOCR.

    Модель — eslav_PP-OCRv5_rec_mobile (PaddleOCR для восточнославянских
    языков) в ONNX Runtime: точность та же, что у PaddleOCR, но ONNX Runtime
    подстраивается под набор инструкций CPU и не падает без AVX2.

    Returns:
        Функция ``(crop_path) -> (текст, признаки неуверенности по строкам)``.
    """
    from rapidocr import RapidOCR
    from rapidocr.utils.typings import LangRec, ModelType, OCRVersion

    engine = RapidOCR(
        params={
            "Rec.lang_type": LangRec.ESLAV,
            "Rec.ocr_version": OCRVersion.PPOCRV5,
            "Rec.model_type": ModelType.MOBILE,
        }
    )

    def read(crop_path: Path) -> tuple[str, tuple[bool, ...]]:
        """Распознаёт фрагмент ячейки."""
        result = engine(str(crop_path))
        items: list[tuple[Box, str, float]] = []
        if result.boxes is not None:
            for box, text, score in zip(result.boxes, result.txts, result.scores):
                rect = (
                    float(box[:, 0].min()),
                    float(box[:, 1].min()),
                    float(box[:, 0].max()),
                    float(box[:, 1].max()),
                )
                items.append((rect, text, float(score)))
        lines = table_structure.join_fragment_lines(items)
        text = "\n".join(line for line, _ in lines)
        return text, tuple(score < LOW_CONFIDENCE_THRESHOLD for _, score in lines)

    return read


def _tesseract_text_predict_fn() -> Callable[[Path], dict]:
    """Создаёт замену PPStructureV3 на Tesseract для текста таблиц.

    Используется, когда PaddlePaddle недоступен. Возвращает JSON той же формы
    (``rec_boxes``/``rec_texts`` в ``table_res_list``), но без ``cell_box_list``:
    области таблиц тогда берутся у TATR. Текст ячеек хуже, чем у PaddleOCR.
    """

    def predict(image_path: Path) -> dict:
        """Распознаёт слова фрагмента таблицы."""
        data = _default_image_data_fn(image_path)
        boxes: list[list[float]] = []
        texts: list[str] = []
        for text, conf, left, top, width, height in zip(
            data["text"], data["conf"], data["left"], data["top"], data["width"], data["height"]
        ):
            word = (text or "").strip()
            if not word or int(conf) < 0:
                continue
            boxes.append([float(left), float(top), float(left + width), float(top + height)])
            texts.append(word)
        return {"table_res_list": [{"table_ocr_pred": {"rec_boxes": boxes, "rec_texts": texts}}]}

    return predict


TEXT_PREDICT_MAX_SIDE = 3000  # PPStructureV3 падает (segfault) на изображениях больше ~3900 px по стороне
CELL_UPSCALE = 2
CELL_PADDING_PX = 12


def _prepare_cell_image(cell: Image.Image) -> Image.Image:
    """Добавляет белое поле и увеличивает фрагмент ячейки.

    Мелкий текст распознаётся точнее, когда знаки крупнее и не прижаты к краю.
    """
    from PIL import ImageOps

    padded = ImageOps.expand(cell, border=CELL_PADDING_PX, fill="white")
    return padded.resize((padded.width * CELL_UPSCALE, padded.height * CELL_UPSCALE), Image.LANCZOS)


def _recognize_cells(
    crop: Image.Image,
    structure: TableStructure,
    texts: list[tuple[Box, str]],
    cell_text_fn: Callable[[Path], tuple[str, tuple[bool, ...]]],
) -> tuple[list[tuple[Box, str]], dict[tuple[int, int], tuple[bool, ...]]]:
    """Распознаёт каждую ячейку с текстом по отдельности.

    Ячейка вырезается по границам своего текста.

    Returns:
        Текстовые блоки ячеек на прежних местах и признаки неуверенности по
        строкам для ячеек с уверенностью ниже ``LOW_CONFIDENCE_THRESHOLD``.
    """
    refined: list[tuple[Box, str]] = []
    low: dict[tuple[int, int], tuple[bool, ...]] = {}
    with tempfile.TemporaryDirectory() as tmp_dir:
        for (r, c), blocks in table_structure.cell_text_boxes(structure, texts).items():
            box = table_structure.text_crop_box(blocks)
            if box is None:
                continue
            cell_path = Path(tmp_dir) / f"cell_{r}_{c}.png"
            _prepare_cell_image(crop.crop(tuple(int(v) for v in box))).save(cell_path)
            text, low_lines = cell_text_fn(cell_path)
            refined.append((box, normalize_header_text(text)))
            if any(low_lines):
                low[(r, c)] = low_lines
    return refined, low


def _box_iou(a: Box, b: Box) -> float:
    """Отношение площади пересечения рамок к площади объединения."""
    ix = max(0.0, min(a[2], b[2]) - max(a[0], b[0]))
    iy = max(0.0, min(a[3], b[3]) - max(a[1], b[1]))
    inter = ix * iy
    union = (a[2] - a[0]) * (a[3] - a[1]) + (b[2] - b[0]) * (b[3] - b[1]) - inter
    return inter / union if union > 0 else 0.0


def _unique_text_blocks(text_json: dict) -> list[tuple[Box, str]]:
    """Убирает повторы текстовых блоков из ответа PaddleOCR.

    На одном фрагменте PaddleOCR возвращает несколько записей таблиц с одним и
    тем же текстом в почти совпадающих рамках. Из блоков с одинаковым текстом
    и IoU ≥ 0.5 остаётся один.
    """
    kept: list[tuple[Box, str]] = []
    for t in text_json.get("table_res_list", []):
        rec = t.get("table_ocr_pred", {})
        for box, text in zip(rec.get("rec_boxes", []), rec.get("rec_texts", [])):
            box = tuple(box)
            if any(k_text == text and _box_iou(k_box, box) >= 0.5 for k_box, k_text in kept):
                continue
            kept.append((box, normalize_header_text(text)))
    return kept


def _merge_overlapping_regions(boxes: list[Box]) -> list[Box]:
    """Объединяет пересекающиеся области (части одной таблицы) в одну."""
    merged = list(boxes)
    changed = True
    while changed:
        changed = False
        for i in range(len(merged)):
            for j in range(i + 1, len(merged)):
                a, b = merged[i], merged[j]
                if _box_iou(a, b) > 0 or (min(a[2], b[2]) > max(a[0], b[0]) and min(a[3], b[3]) > max(a[1], b[1])):
                    merged[i] = (min(a[0], b[0]), min(a[1], b[1]), max(a[2], b[2]), max(a[3], b[3]))
                    del merged[j]
                    changed = True
                    break
            if changed:
                break
    return merged


def paddle_table_regions(page: Image.Image, predict_fn: Callable[[Path], dict]) -> list[Box]:
    """Находит области таблиц через PaddleOCR (PPStructureV3).

    В отличие от TATR, который на странице с несколькими таблицами даёт одну
    общую область, PaddleOCR разделяет таблицы.
    """
    with tempfile.TemporaryDirectory() as tmp_dir:
        page_path = Path(tmp_dir) / "page.png"
        page.save(page_path)
        data = predict_fn(page_path)
    boxes: list[Box] = []
    for t in data.get("table_res_list", []):
        cells = t.get("cell_box_list") or []
        if cells:
            boxes.append((
                min(b[0] for b in cells), min(b[1] for b in cells),
                max(b[2] for b in cells), max(b[3] for b in cells),
            ))
    return _merge_overlapping_regions(boxes)


def table_regions_for_page(page: Image.Image, predict_fn: Callable[[Path], dict]) -> list[Box]:
    """Находит области таблиц на странице.

    Сначала работает лёгкий TATR. Тяжёлый PPStructureV3 (около десятка моделей
    PaddleX) загружается, только если TATR нашёл хотя бы одну область, чтобы
    страницы без таблиц не платили за его загрузку памятью. Найденные области
    уточняет PaddleOCR (он разделяет соседние таблицы); если он ничего не
    нашёл или PaddlePaddle недоступен, остаются области TATR.

    Args:
        page: Изображение страницы.
        predict_fn: Распознавание PPStructureV3.

    Returns:
        Рамки областей таблиц.
    """
    tatr_found = table_structure.detect_table_regions(page)
    if not tatr_found:
        return []
    if not _paddle_runtime_usable():
        return tatr_found
    return paddle_table_regions(page, predict_fn) or tatr_found


class _LazyEngine:
    """Движок, создаваемый при первом вызове и выгружаемый по :meth:`release`.

    Страницы без таблиц-сканов не платят за загрузку моделей, а тяжёлые движки
    не находятся в памяти одновременно.

    Args:
        factory: Создаёт движок — вызываемый объект.
    """

    def __init__(self, factory: Callable[[], Callable]) -> None:
        self._factory = factory
        self._engine: Callable | None = None

    def __call__(self, *args):
        """Вызывает движок, создав его при первом обращении."""
        if self._engine is None:
            self._engine = self._factory()
        return self._engine(*args)

    def release(self) -> None:
        """Выгружает движок."""
        self._engine = None
        gc.collect()


def run_tatr_tables(
    image_path: Path,
    text_predict_fn: Callable[[Path], dict],
    *,
    cell_text_fn: Callable[[Path], str] | None = None,
    region_fn: Callable[[Image.Image], list[Box]] | None = None,
    structure_fn: Callable[[Image.Image], TableStructure] | None = None,
    release_structure: Callable[[], None] | None = None,
    release_cells: Callable[[], None] | None = None,
    img2table_fallback_fn: Callable[[Path], list[tuple[Box, list[TableCell], int, int]]] | None = None,
) -> list[TableBlock]:
    """Распознаёт таблицы страницы через TATR и OCR ячеек.

    Области и структуру (строки, столбцы, объединения) находит TATR, текст —
    OCR по вырезанной с запасом таблице. Пустые столбцы и строки удаляются как
    артефакты детекции. Работа идёт этапами, чтобы тяжёлые движки не были в
    памяти одновременно: сначала области, текст фрагментов и структура всех
    таблиц, затем ``release_structure()``, затем текст ячеек и
    ``release_cells()``.

    Args:
        image_path: Изображение страницы.
        text_predict_fn: Распознавание текста фрагмента таблицы.
        cell_text_fn: Распознавание фрагмента ячейки; None — без повторного
            распознавания по ячейкам.
        region_fn: Поиск областей; по умолчанию
            :func:`npa_ocr.table_structure.detect_table_regions`.
        structure_fn: Поиск структуры; по умолчанию
            :func:`npa_ocr.table_structure.detect_structure`.
        release_structure: Выгрузка движков структуры.
        release_cells: Выгрузка движка ячеек.
        img2table_fallback_fn: Перепроверка подозрительно узких таблиц (см.
            ``_IMG2TABLE_RIGHT_MARGIN_THRESHOLD``); по умолчанию
            :func:`_img2table_tables_for_page`.

    Returns:
        Таблицы страницы.
    """
    structure_fn = structure_fn or table_structure.detect_structure
    img2table_fallback_fn = img2table_fallback_fn or _img2table_tables_for_page
    img2table_candidates: list[tuple[Box, list[TableCell], int, int]] | None = None

    with Image.open(image_path) as opened:
        page = opened.convert("RGB")
    regions = region_fn(page) if region_fn is not None else table_structure.detect_table_regions(page)
    pending: list[tuple] = []
    for x1, y1, x2, y2 in regions:
        crop_box = (
            max(0, int(x1) - _TABLE_CROP_PADDING),
            max(0, int(y1) - _TABLE_CROP_PADDING),
            min(page.width, int(x2) + _TABLE_CROP_PADDING),
            min(page.height, int(y2) + _TABLE_CROP_PADDING),
        )
        crop = page.crop(crop_box)
        scale = min(1.0, TEXT_PREDICT_MAX_SIDE / max(crop.size))
        text_img = crop
        if scale < 1.0:
            text_img = crop.resize((round(crop.width * scale), round(crop.height * scale)), Image.LANCZOS)
        with tempfile.TemporaryDirectory() as tmp_dir:
            crop_path = Path(tmp_dir) / "table_crop.png"
            text_img.save(crop_path)
            text_json = text_predict_fn(crop_path)
        texts = [
            (tuple(v / scale for v in box), text) for box, text in _unique_text_blocks(text_json)
        ]

        structure = structure_fn(crop)
        structure, dropped_cols = table_structure.drop_empty_columns(structure, texts)
        structure = table_structure.expand_rows_to_text(structure, texts)
        structure = table_structure.drop_empty_rows(structure, texts)
        structure = table_structure.split_oversized_rows(structure, texts)
        pending.append((crop_box, (x1, y1, x2, y2), crop, texts, structure, dropped_cols))

    if release_structure is not None:
        release_structure()

    blocks: list[TableBlock] = []
    for crop_box, (x1, y1, x2, y2), crop, texts, structure, dropped_cols in pending:
        low_cells: dict[tuple[int, int], tuple[bool, ...]] = {}
        if cell_text_fn is not None:
            texts, low_cells = _recognize_cells(crop, structure, texts, cell_text_fn)
        structure = table_structure.filter_false_positive_spans(structure, texts)
        cells, n_rows, n_cols = table_structure.build_cells(structure, texts)
        if not cells or not table_structure.is_plausible_table(cells, n_rows, n_cols):
            continue
        for cell in cells:
            cell.low_lines = low_cells.get((cell.row, cell.col), ())
        content = [box for (_r, _c), blk in table_structure.cell_text_boxes(structure, texts).items() for box, _ in blk]
        if content:
            bx1 = min(b[0] for b in content) + crop_box[0]
            by1 = min(b[1] for b in content) + crop_box[1]
            bx2 = max(b[2] for b in content) + crop_box[0]
            by2 = max(b[3] for b in content) + crop_box[1]
        else:
            bx1, by1, bx2, by2 = x1, y1, x2, y2

        right_margin = (page.width - bx2) / page.width if page.width else 0.0
        if right_margin > _IMG2TABLE_RIGHT_MARGIN_THRESHOLD:
            if img2table_candidates is None:
                img2table_candidates = img2table_fallback_fn(image_path)
            match = _best_img2table_match(by1, by2, img2table_candidates)
            if match is not None and (match[0][2] - match[0][0]) > (bx2 - bx1):
                (bx1, by1, bx2, by2), cells, n_rows, n_cols = match
                dropped_cols = 0

        blocks.append(
            TableBlock(
                y=by1,
                bbox=(bx1, by1, bx2, by2),
                cells=cells,
                n_rows=n_rows,
                n_cols=n_cols,
                dropped_columns=dropped_cols,
            )
        )
    if not regions:
        # TATR не нашёл ни одной области: короткий хвост многостраничной таблицы он не
        # замечает ни на каком безопасном пороге. Страница проходится img2table; его ложные
        # таблицы отсеивает is_plausible_table. Области, найденные TATR и отклонённые как
        # неправдоподобные, сюда не попадают.
        if img2table_candidates is None:
            img2table_candidates = img2table_fallback_fn(image_path)
        for bbox, cells, n_rows, n_cols in img2table_candidates:
            blocks.append(
                TableBlock(y=bbox[1], bbox=bbox, cells=cells, n_rows=n_rows, n_cols=n_cols, dropped_columns=0)
            )

    if release_cells is not None:
        release_cells()
    return blocks


def _load_tatr_engines() -> tuple["_LazyEngine", "_LazyEngine"]:
    """Готовит движки для ``table_engine="tatr"``.

    TATR загружается раньше PaddleOCR: на Windows torch нужно импортировать
    раньше paddle. Остальные движки создаются лениво. Без PaddlePaddle (CPU без
    AVX2) текст фрагментов таблиц читает Tesseract.

    Returns:
        Движок текста фрагментов таблиц и движок текста ячеек.
    """
    table_structure.load_tatr()
    cell_text_fn = _LazyEngine(_default_cell_text_fn)
    if _paddle_runtime_usable():
        return _LazyEngine(_default_predict_tables_fn), cell_text_fn
    print("PaddleOCR недоступен на этом CPU (нет AVX2 / SIGILL) — регионы таблиц через Tesseract", file=sys.stderr)
    return _LazyEngine(_tesseract_text_predict_fn), cell_text_fn


def run_vlm_table_blocks(
    image_path: Path,
    chat_fn: Callable[[Image.Image], str],
    *,
    ask_fn: "vlm_reread.AskFn | None" = None,
    timer: "StepTimer | None" = None,
    page: int | None = None,
) -> list[TableBlock]:
    """Распознаёт таблицы страницы через DeepSeek-OCR и сетку по линиям.

    Ячейки с текстом только от Tesseract, поставленным по положению или
    исправленным повторным чтением выделяются для сверки со сканом.

    Args:
        image_path: Изображение страницы.
        chat_fn: Запрос к DeepSeek-OCR по странице.
        ask_fn: Запрос к модели для повторного чтения ячеек; None — без него.
        timer: Учёт времени шагов.
        page: Номер страницы для учёта времени.

    Returns:
        Таблицы страницы.

    Raises:
        VlmUnavailable: vLLM недоступен; запасной путь выбирает вызывающий код.
    """
    chat_seconds = [0.0]
    reread_seconds = [0.0]

    def timed_chat(img: Image.Image) -> str:
        """``chat_fn`` с учётом времени."""
        t0 = time.perf_counter()
        try:
            return chat_fn(img)
        finally:
            chat_seconds[0] = time.perf_counter() - t0

    def timed_reread(crop: Image.Image, doubt: list[dict]) -> dict[tuple[int, int], str]:
        """Повторное чтение ячеек с учётом времени."""
        t0 = time.perf_counter()
        try:
            return vlm_reread.reread_cells(crop, doubt, ask_fn)
        finally:
            reread_seconds[0] += time.perf_counter() - t0

    t0 = time.perf_counter()
    tables = vlm_tables.run_vlm_tables(
        image_path, chat_fn=timed_chat, reread_fn=timed_reread if ask_fn is not None else None
    )
    elapsed = time.perf_counter() - t0
    if timer is not None:
        timer.add(page, "таблицы: DeepSeek-OCR (vLLM)", chat_seconds[0])
        timer.add(page, "таблицы: сетка по линиям", elapsed - chat_seconds[0] - reread_seconds[0])
        if reread_seconds[0]:
            timer.add(page, "таблицы: дочитка ячеек (vLLM)", reread_seconds[0])
    return [
        TableBlock(
            y=float(t.bbox[1] + t.info["ye"][0]) if t.src == "grid" and t.info.get("ye") else float(t.bbox[1]),
            # Поле вокруг кропа нужно OCR, но не принадлежит таблице.
            # Исключение этого поля из обычного текста съедало последнюю
            # строку заголовков, расположенных вплотную к верхней рамке.
            bbox=(float(t.bbox[0] + t.info["xe"][0]), float(t.bbox[1] + t.info["ye"][0]),
                  float(t.bbox[0] + t.info["xe"][-1]), float(t.bbox[1] + t.info["ye"][-1]))
                 if t.src == "grid" and t.info.get("xe") and t.info.get("ye") else tuple(map(float, t.bbox)),
            cells=t.cells,
            n_rows=t.n_rows,
            n_cols=t.n_cols,
            column_edges=[(x - t.info["xe"][0]) / max(1, t.info["xe"][-1] - t.info["xe"][0])
                          for x in t.info.get("xe", [])] if t.src == "grid" else [],
            source_pages=[page] if page is not None else [],
        )
        for t in tables
        if t.cells
    ]


def ocr_page(
    image_path: Path,
    *,
    predict_tables_fn: Callable[[Path], dict] | None = None,
    image_data_fn: Callable[[Path], dict],
    confidence_threshold: int = DEFAULT_CONFIDENCE_THRESHOLD,
    authority_name: str | None = None,
    table_blocks_fn: Callable[[Path], list[TableBlock]] | None = None,
    detect_marks_fn: Callable[[Path, list[Box]], list[mark_detection.Mark]] | None = None,
    recover_covered_text_fn: Callable[[Path], str] | None = None,
    clean_marks_fn: Callable[[list[list[vlm_marks.Word]]], list[list[vlm_marks.Word]]] | None = None,
    reread_text_fn: Callable[[list[list[vlm_marks.Word]]],
                             tuple[list[list[vlm_marks.Word]], list[vlm_text.Formula]]] | None = None,
) -> list[PageBlock]:
    """Распознаёт страницу: таблицы и текст в порядке чтения.

    Таблицы распознаются раньше текста, поэтому ``clean_marks_fn`` и
    ``reread_text_fn`` могут использовать ответ DeepSeek на страницу.

    Args:
        image_path: Изображение страницы.
        predict_tables_fn: Таблицы через PaddleOCR; используется, если не
            задан ``table_blocks_fn``.
        image_data_fn: Распознавание слов Tesseract.
        confidence_threshold: Порог уверенности Tesseract.
        authority_name: Наименование органа.
        table_blocks_fn: Распознавание таблиц страницы.
        detect_marks_fn: См. :func:`run_tesseract_text`.
        recover_covered_text_fn: См. :func:`run_tesseract_text`.
        clean_marks_fn: См. :func:`run_tesseract_text`.
        reread_text_fn: См. :func:`run_tesseract_text`.

    Returns:
        Блоки страницы сверху вниз.
    """
    if table_blocks_fn is not None:
        tables = table_blocks_fn(image_path)
    else:
        tables = run_paddle_tables(image_path, predict_tables_fn)
    text_blocks = run_tesseract_text(
        image_path,
        [t.bbox for t in tables],
        image_data_fn=image_data_fn,
        confidence_threshold=confidence_threshold,
        authority_name=authority_name,
        detect_marks_fn=detect_marks_fn,
        recover_covered_text_fn=recover_covered_text_fn,
        clean_marks_fn=clean_marks_fn,
        reread_text_fn=reread_text_fn,
    )
    if tables and table_blocks_fn is not None:
        from npa_ocr.coverage import absorb_ruled_notes, recover_table_closers
        with Image.open(image_path) as source_image:
            tables, text_blocks = absorb_ruled_notes(source_image, tables, text_blocks)
            text_blocks = recover_table_closers(source_image, tables, text_blocks)
    blocks: list[PageBlock] = [*tables, *text_blocks]
    blocks.sort(key=lambda b: b.y)
    return blocks


# --- Сборка DOCX ---


def _add_table(doc, block: TableBlock, *, diagnostic: bool = False) -> None:
    """Добавляет таблицу в документ с объединением ячеек.

    Объединения разных ячеек могут конфликтовать и давать непрямоугольную
    область, которую python-docx не объединяет (``InvalidSpanError``). Такая
    ячейка остаётся необъединённой: текст не теряется, документ собирается.
    Неуверенные строки ячеек выделяются жёлтым.
    """
    from docx.enum.text import WD_COLOR_INDEX
    from docx.exceptions import InvalidSpanError

    if diagnostic and block.source_pages:
        pages = ", ".join(map(str, block.source_pages))
        doc.add_paragraph(f"Источник: PDF, страницы {pages}.", style="Caption")
    table = doc.add_table(rows=block.n_rows, cols=block.n_cols)
    table.style = "Table Grid"
    if len(block.column_edges) == block.n_cols + 1:
        # Продолжения могут иметь дополнительные границы объединений.
        # Их узкие промежутки сохраняют ширину скана, а не получают
        # одинаковую ширину с полноценными графами.
        section = doc.sections[-1]
        available = int(section.page_width - section.left_margin - section.right_margin)
        widths = [int(available * (right - left))
                  for left, right in zip(block.column_edges, block.column_edges[1:])]
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = width
        for i, target in enumerate(table._cells):
            target.width = widths[i % block.n_cols]
    merged: set[tuple[int, int]] = set()
    for cell in block.cells:
        if (cell.row, cell.col) in merged:
            continue
        target = table.cell(cell.row, cell.col)
        if cell.row_span > 1 or cell.col_span > 1:
            bottom_right = table.cell(
                cell.row + cell.row_span - 1, cell.col + cell.col_span - 1
            )
            try:
                target = target.merge(bottom_right)
            except InvalidSpanError:
                target = table.cell(cell.row, cell.col)
            else:
                for r in range(cell.row, cell.row + cell.row_span):
                    for c in range(cell.col, cell.col + cell.col_span):
                        merged.add((r, c))
        lines = cell.text.split("\n")
        if any(cell.low_lines) and len(lines) == len(cell.low_lines):
            target.text = ""
            paragraph = target.paragraphs[0]
            for i, (line, low) in enumerate(zip(lines, cell.low_lines)):
                if i:
                    paragraph.add_run().add_break()
                run = paragraph.add_run(line)
                if low:
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            target.text = cell.text
    from npa_ocr.docx_layout import format_table
    format_table(table, block)


_INDENT_MIN_PX = 40  # около 5 мм при 200 DPI — отступ первой строки абзаца
_SHORT_LINE_RATIO = 0.12  # строка, не доходящая до правого края на эту долю ширины, — конец абзаца
_CENTER_TOLERANCE_RATIO = 0.015
_CENTER_MAX_WIDTH_RATIO = 0.95
_PARAGRAPH_GAP_RATIO = 1.5  # зазор больше этой доли высоты строки — новый абзац


Segment = tuple[str, bool, str]  # слово или пробел, низкая уверенность, разметка (TextBlock.rich_words)


def _block_segments(b: TextBlock) -> list[Segment]:
    """Слова строки и пробелы между ними с признаками неуверенности и разметкой."""
    words = b.text.split(" ")
    flags = b.low_words if len(b.low_words) == len(words) else (False,) * len(words)
    rich = b.rich_words if len(b.rich_words) == len(words) else ("",) * len(words)
    segments: list[Segment] = []
    for i, (word, low, markup) in enumerate(zip(words, flags, rich)):
        if i:
            segments.append((" ", False, ""))
        segments.append((word, low, markup))
    return segments


def _paragraph_segments(group: list[TextBlock]) -> list[Segment]:
    """Сегменты строк абзаца, соединённых пробелом."""
    segments: list[Segment] = []
    for i, b in enumerate(group):
        if i:
            segments.append((" ", False, ""))
        segments.extend(_block_segments(b))
    return segments


def _text_paragraphs(run: list[TextBlock], dpi: int) -> list[tuple[str, str, float, list[Segment]]]:
    """Склеивает строки непрерывного участка текста в абзацы по геометрии.

    Новый абзац начинается после центрированной строки, с отступа первой
    строки, после короткой строки и после большого зазора. Начертание и сноски
    не определяются.

    Args:
        run: Строки участка сверху вниз.
        dpi: Разрешение страницы для перевода отступа в пункты.

    Returns:
        Абзацы: текст, выравнивание (``"left"`` или ``"center"``), отступ первой
        строки в пунктах и сегменты.
    """
    from statistics import median

    if not run:
        return []
    # A recovered number can be a separate OCR line slightly above the
    # date/city baseline. Rejoin this one spatially confirmed requisites row.
    run=list(run)
    for number in list(run):
        if not re.fullmatch(r"№\s*\d{1,5}",number.text):
            continue
        date=next((b for b in run if b is not number
                   and re.fullmatch(r"\d{2}\.\d{2}\.\d{4}\s+[А-Яа-яЁё. -]+",b.text)
                   and b.x1<number.x0 and max(b.y,number.y)<min(b.y+b.height,number.y+number.height)),None)
        if date is not None:
            flags=tuple(date.low_words)+tuple(number.low_words)
            run[run.index(date)]=replace(date,text=date.text+' '+number.text,x1=number.x1,
                                         low_words=flags)
            run.remove(number)
    if not all(b.x1 > b.x0 for b in run):
        return [(b.text, "left", 0.0, _paragraph_segments([b])) for b in run]

    area_left = min(b.x0 for b in run)
    area_right = max(b.x1 for b in run)
    area_width = area_right - area_left
    area_center = (area_left + area_right) / 2
    line_height = median(b.height for b in run if b.height > 0) if any(b.height > 0 for b in run) else 0.0

    def is_centered(b: TextBlock) -> bool:
        """Строка выровнена по центру участка."""
        width = b.x1 - b.x0
        return (
            abs((b.x0 + b.x1) / 2 - area_center) <= _CENTER_TOLERANCE_RATIO * area_width
            and width <= _CENTER_MAX_WIDTH_RATIO * area_width
        )

    def px_to_pt(px: float) -> float:
        """Пиксели в пункты."""
        return px * 72 / dpi

    paragraphs: list[list[TextBlock]] = []
    prev: TextBlock | None = None
    for b in run:
        centered = is_centered(b)
        if prev is None:
            new_paragraph = True
        elif re.fullmatch(r'\s*»\s*\.?\s*',b.text) or re.fullmatch(r'\s*»\s*\.?\s*',prev.text):
            new_paragraph = True
        elif centered or is_centered(prev):
            new_paragraph = not (centered and is_centered(prev) and line_height
                                 and b.y - (prev.y + prev.height) <= line_height
                                 and not re.search(r"[.!?;:]\s*$", prev.text))
        elif b.x0 - area_left > _INDENT_MIN_PX:
            new_paragraph = True
        elif (len(b.text)>3 and len(prev.text)>3 and abs(b.x0-prev.x0)<20
              and max(b.x1-b.x0,prev.x1-prev.x0)<.65*area_width
              and line_height and b.y-(prev.y+prev.height)<=1.5*line_height
              and not re.search(r"[.!?;:]\s*$",prev.text)
              and not re.match(r"\s*\d+[.)]\s",b.text)):
            new_paragraph = False
        elif prev.x1 < area_right - _SHORT_LINE_RATIO * area_width:
            new_paragraph = True
        elif line_height and b.y - (prev.y + prev.height) > _PARAGRAPH_GAP_RATIO * line_height:
            new_paragraph = True
        else:
            new_paragraph = False
        if new_paragraph:
            paragraphs.append([b])
        else:
            paragraphs[-1].append(b)
        prev = b

    result: list[tuple[str, str, float, list[Segment]]] = []
    for group in paragraphs:
        first = group[0]
        text = " ".join(b.text for b in group)
        segments = _paragraph_segments(group)
        if is_centered(first):
            result.append((text, "center", 0.0, segments))
        else:
            indent = first.x0 - area_left
            first_line = px_to_pt(indent) if indent > _INDENT_MIN_PX else 0.0
            result.append((text, "left", first_line, segments))
    return result


_HEADER_FOOTER_BAND_RATIO = 0.06
_PAGE_NUMBER_RE = re.compile(r"^[\s\-–—.]*\d{1,4}[\s\-–—.]*$")


def _normalize_for_repeat(text: str) -> str:
    """Ключ для поиска повторов: нижний регистр, цифры заменены на «#»."""
    return re.sub(r"\d", "#", text.strip().lower())


def remove_running_headers_and_footers(pages: list[PageResult]) -> list[PageResult]:
    """Убирает колонтитулы со страниц, распознанных OCR.

    Колонтитул — строка в верхней или нижней полосе страницы (6 % высоты),
    которая состоит из номера страницы или повторяется (с точностью до цифр)
    на половине страниц и более. Таблицы и строки вне полос не меняются.

    Args:
        pages: Страницы документа.

    Returns:
        Страницы без колонтитулов.
    """
    ocr_pages = [p for p in pages if p.ocr_applied and p.page_height > 0]
    if not ocr_pages:
        return pages

    def in_band(block: TextBlock, height: float) -> bool:
        """Строка в полосе колонтитула."""
        top_limit = _HEADER_FOOTER_BAND_RATIO * height
        bottom_limit = (1 - _HEADER_FOOTER_BAND_RATIO) * height
        return block.y < top_limit or block.y + block.height > bottom_limit

    repeats: dict[str, int] = {}
    for page in ocr_pages:
        seen = {
            _normalize_for_repeat(b.text)
            for b in page.blocks
            if isinstance(b, TextBlock) and in_band(b, page.page_height)
        }
        for key in seen:
            repeats[key] = repeats.get(key, 0) + 1
    threshold = max(2, len(ocr_pages) // 2)

    result: list[PageResult] = []
    for page in pages:
        if not page.ocr_applied or page.page_height <= 0:
            result.append(page)
            continue
        kept = []
        for block in page.blocks:
            if isinstance(block, TextBlock) and in_band(block, page.page_height):
                key = _normalize_for_repeat(block.text)
                if _PAGE_NUMBER_RE.match(block.text.strip()) or repeats.get(key, 0) >= threshold:
                    continue
            kept.append(block)
        result.append(PageResult(
            page_number=page.page_number,
            ocr_applied=page.ocr_applied,
            blocks=kept,
            plain_text=page.plain_text,
            page_height=page.page_height,
            review_issues=page.review_issues,
        ))
    return result


# Продолжение таблицы: верхняя часть доходит ниже этой доли высоты страницы, продолжение
# начинается выше этой (над ним — только колонтитул).
_STITCH_BOTTOM_RATIO = 0.80
_STITCH_TOP_RATIO = 0.15


def stitch_tables(pages: list[PageResult]) -> tuple[list[PageResult], list[str]]:
    """Сшивает таблицы, продолжающиеся на следующей странице.

    Продолжение — первая таблица страницы, если она начинается сверху,
    последняя таблица предыдущей страницы доходит до низа и число столбцов
    совпадает. Продолжение дописывается в таблицу предыдущей страницы
    (:func:`npa_ocr.table_stitch.merge_continuation`) и убирается со своей;
    цепочка может идти через много страниц.

    Args:
        pages: Страницы документа; не изменяются.

    Returns:
        Новые страницы и журнал склеек.
    """
    result: list[PageResult] = []
    log: list[str] = []
    # последняя таблица: (страница в result, индекс блока, номер страницы последней части, низ части)
    last: tuple[PageResult, int, int, float] | None = None
    for page in pages:
        if not page.ocr_applied or page.page_height <= 0:
            result.append(page)
            last = None
            continue
        page = replace(page, blocks=list(page.blocks))
        result.append(page)
        if last is not None and last[2] == page.page_number - 1 and last[3] >= _STITCH_BOTTOM_RATIO:
            host_page, index, _, _ = last
            host = host_page.blocks[index]
            note = next((c for c in host.cells if c.col == 0 and c.col_span == host.n_cols
                         and c.row + c.row_span == host.n_rows and len(c.text) > 60
                         and re.search(r"ГОСТ|регистрацион|Примеч", c.text, re.I)), None)
            fragments: list[TextBlock] = []
            closed = False
            for candidate in page.blocks:
                if not isinstance(candidate, TextBlock) or candidate.y > .25 * page.page_height:
                    break
                if re.fullmatch(r'\s*[»”"]\s*\.?\s*', candidate.text):
                    closed = bool(fragments)
                    break
                if not fragments and (candidate.y > .15 * page.page_height
                                      or not candidate.text[:1].islower()):
                    break
                if re.match(r"\s*\d+[.)]\s", candidate.text):
                    break
                fragments.append(candidate)
            if note is not None and closed:
                original_lines = note.text.split("\n")
                low = (note.low_lines if len(note.low_lines) == len(original_lines)
                       else [any(note.low_lines)] * len(original_lines))
                joined = replace(note, text=note.text + "\n" + "\n".join(b.text for b in fragments),
                                 low_lines=list(low) + [True] * len(fragments))
                host_page.blocks[index] = replace(host, cells=[joined if c is note else c for c in host.cells],
                    source_pages=sorted(set(host.source_pages + [page.page_number])))
                page.blocks = page.blocks[len(fragments):]
                page.continues_table = True
                log.append(f"стр. {page.page_number}: продолжение примечания возвращено в ячейку таблицы")
                last = None
        kept: list[PageBlock] = []
        first_table = True
        for block in page.blocks:
            if not isinstance(block, TableBlock):
                kept.append(block)
                continue
            top_ratio = block.bbox[1] / page.page_height
            bottom_ratio = block.bbox[3] / page.page_height
            host = last[0].blocks[last[1]] if last is not None else None
            aligned = (table_stitch.align_columns(host.cells, host.column_edges, block.cells, block.column_edges)
                       if host is not None and host.column_edges and block.column_edges else None)
            if (
                first_table
                and last is not None
                and last[2] == page.page_number - 1
                and last[3] >= _STITCH_BOTTOM_RATIO
                and top_ratio <= _STITCH_TOP_RATIO
                and (aligned is not None or (not (host.column_edges and block.column_edges)
                                             and host.n_cols == block.n_cols))
            ):
                host_page, index, _, _ = last
                host = host_page.blocks[index]
                top_cells, cont_cells, edges = aligned if aligned is not None else (host.cells, block.cells, [])
                cells, n_rows, done = table_stitch.merge_continuation(
                    top_cells, host.n_rows, cont_cells, block.n_rows
                )
                host_page.blocks[index] = replace(host, cells=cells, n_rows=n_rows,
                    n_cols=len(edges) - 1 if edges else host.n_cols,
                    column_edges=edges or host.column_edges,
                    source_pages=sorted(set(host.source_pages + block.source_pages)))
                page.continues_table = True
                log.append(
                    f"стр. {page.page_number}: продолжение таблицы со стр. {host_page.page_number} — "
                    + ("; ".join(done) or "склеено без правок")
                )
                last = (host_page, index, page.page_number, bottom_ratio)
            else:
                kept.append(block)
                last = (page, len(kept) - 1, page.page_number, bottom_ratio)
            first_table = False
        page.blocks[:] = kept
    return result, log


def build_docx(pages: list[PageResult], out_path: Path, *, dpi: int = DEFAULT_DPI,
               diagnostic: bool = False) -> None:
    """Собирает DOCX из распознанных страниц.

    Страницы с текстовым слоем пишутся как есть. На страницах OCR строки
    склеиваются в абзацы по геометрии (:func:`_text_paragraphs`), таблицы
    становятся таблицами Word с объединениями, формулы — формулами Word.
    Неуверенно распознанный текст выделяется жёлтым.

    Args:
        pages: Страницы документа.
        out_path: Путь к DOCX.
        dpi: Разрешение страниц для перевода отступов в пункты.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.shared import Pt, Cm

    doc = Document()
    from npa_ocr.docx_layout import set_document_language, keep_table_titles, format_requisites
    set_document_language(doc)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1)
    section.top_margin = section.bottom_margin = Cm(1)
    section.header_distance = section.footer_distance = Cm(.5)
    has_low_confidence = any(
        (isinstance(block, TableBlock) and any(any(cell.low_lines) for cell in block.cells))
        or (isinstance(block, TextBlock) and any(block.low_words))
        for page in pages
        if page.ocr_applied
        for block in page.blocks
    )
    if has_low_confidence:
        doc.add_paragraph(
            "Жёлтым выделен неуверенно распознанный или восстановленный текст: проверьте его по скану."
        )
    for i, page in enumerate(pages):
        # A source page boundary is not a layout boundary after reflow.
        # Keep it only in diagnostic exports and for untouched text layers.
        if i > 0 and not page.continues_table and (diagnostic or not page.ocr_applied):
            doc.add_page_break()
        if not page.ocr_applied and not page.blocks:
            doc.add_paragraph(page.plain_text or "")
            continue
        run: list[TextBlock] = []

        def flush_run() -> None:
            """Пишет накопленные строки абзацами."""
            for _text, align, first_line, segments in _text_paragraphs(run, dpi):
                para = doc.add_paragraph()
                for segment, low, markup in segments:
                    if markup:
                        docx_math.add_rich(para, markup, low)
                        continue
                    word_run = para.add_run(segment)
                    if low:
                        word_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                if align == "center":
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif first_line:
                    para.paragraph_format.first_line_indent = Pt(first_line)
                format_requisites(para,section.page_width-section.left_margin-section.right_margin)
            run.clear()

        for block in page.blocks:
            if isinstance(block, TextBlock):
                run.append(block)
            elif isinstance(block, FormulaBlock):
                flush_run()
                docx_math.add_formula(doc, block.latex)
            else:
                flush_run()
                _add_table(doc, block, diagnostic=diagnostic)
        flush_run()
    keep_table_titles(doc)
    doc.save(str(out_path))


# --- Движки по умолчанию (тяжёлые импорты — только здесь) ---


def _default_predict_tables_fn() -> Callable[[Path], dict]:
    """Создаёт распознаватель таблиц PPStructureV3.

    Модели загружаются один раз на документ. oneDNN (``enable_mkldnn``)
    отключён: с ним движок падает на CPU без AVX2. Распознавание формул,
    диаграмм и печатей отключено как ненужное.

    Returns:
        Функция ``(image_path) -> dict`` в формате ``save_to_json``.
    """
    os.environ.setdefault("FLAGS_use_mkldnn", "false")
    from paddleocr import PPStructureV3

    pipeline = PPStructureV3(
        lang="ru",
        use_doc_orientation_classify=True,
        use_doc_unwarping=False,
        use_textline_orientation=True,
        use_formula_recognition=False,
        use_chart_recognition=False,
        use_seal_recognition=False,
        enable_mkldnn=False,
    )

    def predict(image_path: Path) -> dict:
        """Распознаёт таблицы изображения."""
        # use_table_orientation_classify=False: иначе PPStructureV3 загружает ещё один
        # набор моделей детекции и распознавания (около четырёх) для проверки поворота
        # каждой таблицы, хотя поворот страницы уже проверен. Результат на таблицах без
        # собственного поворота тот же; поворот отдельной таблицы не определяется.
        results = list(pipeline.predict(str(image_path), use_table_orientation_classify=False))
        # JSON — через документированный save_to_json, а не через внутренние поля результата
        with tempfile.TemporaryDirectory() as tmp:
            results[0].save_to_json(tmp)
            json_path = next(Path(tmp).glob("*.json"))
            with open(json_path, encoding="utf-8") as f:
                return json.load(f)

    return predict


# Установщик winget (tesseract-ocr.tesseract) кладёт программу сюда, но не всегда
# добавляет её в PATH текущего процесса.
_WINDOWS_TESSERACT_DEFAULT = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


def _resolve_tesseract_cmd() -> str | None:
    """Путь к программе Tesseract.

    Берётся из ``TESSERACT_CMD``, иначе — стандартный путь установки на
    Windows, если Tesseract нет в PATH.

    Returns:
        Путь или None, если ``pytesseract`` найдёт программу сам.
    """
    import shutil

    env_cmd = os.environ.get("TESSERACT_CMD")
    if env_cmd:
        return env_cmd
    if shutil.which("tesseract"):
        return None  # уже на PATH, pytesseract найдёт сам
    if Path(_WINDOWS_TESSERACT_DEFAULT).is_file():
        return _WINDOWS_TESSERACT_DEFAULT
    return None


# Автоматическая сегментация Tesseract (psm=3) иногда не находит блок текста, отделённый
# от основного большим пустым полем, например строку подписи: строка пропадает целиком.
# Если под последним найденным словом пусто больше этой доли высоты страницы, нижняя
# полоса читается повторно. У обычных страниц пустой остаток 0.06–0.10, у страниц с
# пропущенной подписью — около 0.16.
_BOTTOM_GAP_RATIO = 0.08
_BOTTOM_GAP_MARGIN_PX = 20  # отступ от последнего найденного слова перед повторным кропом
_BOTTOM_GAP_MIN_WORD_LEN = 2  # см. _recover_bottom_gap_text


def _recover_bottom_gap_text(img: "Image.Image", data: dict) -> None:
    """Дочитывает нижнюю полосу страницы, пропущенную основным проходом.

    Полоса под последним найденным словом (см. ``_BOTTOM_GAP_RATIO``)
    читается в режиме psm=6 («единый блок текста»). Найденные слова
    добавляются, только если среди них есть уверенное слово не короче
    ``_BOTTOM_GAP_MIN_WORD_LEN`` букв и цифр: иначе это шум скана, росчерк
    подписи или линия таблицы, которые psm=6 читает как короткое «слово».
    Добавляются все слова полосы, неуверенные выделяются как обычно.

    Args:
        img: Изображение страницы.
        data: Результат ``pytesseract.image_to_data``; дополняется на месте.
    """
    import re

    import pytesseract

    bottoms = [
        data["top"][i] + data["height"][i]
        for i in range(len(data["text"]))
        if data["text"][i].strip()
    ]
    max_bottom = max(bottoms, default=0)
    if img.height <= 0:
        return
    gap_ratio = (img.height - max_bottom) / img.height
    crop_y0 = max_bottom + _BOTTOM_GAP_MARGIN_PX
    if gap_ratio <= _BOTTOM_GAP_RATIO or crop_y0 >= img.height:
        return

    crop = img.crop((0, crop_y0, img.width, img.height))
    extra = pytesseract.image_to_data(crop, lang="rus", config="--psm 6", output_type=pytesseract.Output.DICT)
    has_confident_word = any(
        extra["text"][i].strip()
        and extra["conf"][i] >= DEFAULT_CONFIDENCE_THRESHOLD
        and len(re.sub(r"\W", "", extra["text"][i], flags=re.UNICODE)) >= _BOTTOM_GAP_MIN_WORD_LEN
        for i in range(len(extra["text"]))
    )
    if not has_confident_word:
        return

    block_offset = (max(data["block_num"], default=0) if data["block_num"] else 0) + 1
    for i in range(len(extra["text"])):
        if not extra["text"][i].strip():
            continue
        data["level"].append(extra["level"][i])
        data["page_num"].append(extra["page_num"][i])
        data["block_num"].append(extra["block_num"][i] + block_offset)
        data["par_num"].append(extra["par_num"][i])
        data["line_num"].append(extra["line_num"][i])
        data["word_num"].append(extra["word_num"][i])
        data["left"].append(extra["left"][i])
        data["top"].append(extra["top"][i] + crop_y0)
        data["width"].append(extra["width"][i])
        data["height"].append(extra["height"][i])
        data["conf"].append(extra["conf"][i])
        data["text"].append(extra["text"][i])


def _default_image_data_fn(image_path: Path) -> dict:
    """Распознаёт слова страницы Tesseract (русский язык).

    Returns:
        Результат ``pytesseract.image_to_data`` в виде словаря, с дочитанной
        нижней полосой (:func:`_recover_bottom_gap_text`).
    """
    import pytesseract
    from PIL import Image

    tesseract_cmd = _resolve_tesseract_cmd()
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
    # файл закрывается явно: на Windows открытый дескриптор не даёт удалить временный PNG
    with Image.open(image_path) as img:
        img.load()
        data = pytesseract.image_to_data(img, lang="rus", output_type=pytesseract.Output.DICT)
        _recover_bottom_gap_text(img, data)
        return data


# --- Поворот страницы ---

# Альбомные таблицы в книжном скане лежат на боку, и ни Tesseract, ни модели таблиц их так не
# читают. Угол даёт Tesseract OSD («Rotate: 90»); при меньшей уверенности (мало текста на
# странице) страница считается ровной.
OSD_MIN_CONFIDENCE = 2.0


def parse_osd_rotation(osd: str) -> int:
    """Угол поворота страницы по выводу Tesseract OSD.

    Args:
        osd: Вывод ``pytesseract.image_to_osd``.

    Returns:
        Угол по часовой стрелке (0, 90, 180 или 270); 0 при уверенности ниже
        ``OSD_MIN_CONFIDENCE``.
    """
    rot = re.search(r"Rotate: (\d+)", osd)
    conf = re.search(r"Orientation confidence: ([\d.]+)", osd)
    if not rot or not conf or float(conf.group(1)) < OSD_MIN_CONFIDENCE:
        return 0
    return int(rot.group(1)) % 360


def _default_orientation_fn(image_path: Path) -> int:
    """Определяет поворот страницы через Tesseract OSD.

    На странице почти без текста OSD завершается ошибкой; такая страница
    считается ровной.

    Returns:
        Угол по часовой стрелке.
    """
    import pytesseract

    tesseract_cmd = _resolve_tesseract_cmd()
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
    try:
        with Image.open(image_path) as img:
            img.load()
            return parse_osd_rotation(pytesseract.image_to_osd(img))
    except (pytesseract.TesseractError, pytesseract.TesseractNotFoundError):
        return 0


def rotate_page_image(image_path: Path, rotation: int) -> tuple[int, int]:
    """Поворачивает PNG страницы по часовой стрелке и перезаписывает файл.

    Returns:
        Новые ширина и высота.
    """
    with Image.open(image_path) as img:
        rotated = img.rotate(-rotation, expand=True)
    rotated.save(image_path)
    return rotated.size


def _default_detect_marks_fn(image_path: Path, low_word_boxes: list[Box]) -> list[mark_detection.Mark]:
    """Ищет метки с классификацией через Ollama (:func:`npa_ocr.ai_escalation.classify_mark`)."""
    return mark_detection.detect_marks(image_path, low_word_boxes, classify_fn=ai_escalation.classify_mark)


# --- Постраничное сохранение и возобновление ---

# Сырые результаты страницы (блоки с координатами, таблицы, флаги уверенности) пишутся
# на диск сразу после распознавания страницы. Сборка DOCX читает их, поэтому после
# остановки (нехватка памяти, сбой, перезагрузка) уже распознанные страницы не
# пересчитываются. Колонтитулы и абзацы собираются только при сборке, по всем страницам.
PAGE_CACHE_FORMAT = 1


def _block_to_dict(block: PageBlock) -> dict:
    """Переводит блок страницы в JSON-совместимый словарь."""
    if isinstance(block, TableBlock):
        return {
            "type": "table",
            "y": block.y,
            "bbox": list(block.bbox),
            "n_rows": block.n_rows,
            "n_cols": block.n_cols,
            "dropped_columns": block.dropped_columns,
            "column_edges": block.column_edges,
            "source_pages": block.source_pages,
            "cells": [{**asdict(c), "low_lines": list(c.low_lines)} for c in block.cells],
        }
    if isinstance(block, FormulaBlock):
        return {"type": "formula", "y": block.y, "latex": block.latex, "bbox": list(block.bbox)}
    return {
        "type": "text",
        "y": block.y,
        "text": block.text,
        "x0": block.x0,
        "x1": block.x1,
        "height": block.height,
        "low_words": list(block.low_words),
        "rich_words": list(block.rich_words),
    }


def _block_from_dict(data: dict) -> PageBlock:
    """Восстанавливает блок страницы из словаря :func:`_block_to_dict`."""
    if data["type"] == "table":
        return TableBlock(
            y=data["y"],
            bbox=tuple(data["bbox"]),
            cells=[TableCell(**{**c, "low_lines": tuple(c["low_lines"])}) for c in data["cells"]],
            n_rows=data["n_rows"],
            n_cols=data["n_cols"],
            dropped_columns=data["dropped_columns"],
            column_edges=data.get("column_edges", []),
            source_pages=data.get("source_pages", []),
        )
    if data["type"] == "formula":
        return FormulaBlock(y=data["y"], latex=data["latex"], bbox=tuple(data["bbox"]))
    return TextBlock(
        y=data["y"],
        text=data["text"],
        x0=data["x0"],
        x1=data["x1"],
        height=data["height"],
        low_words=tuple(data["low_words"]),
        rich_words=tuple(data.get("rich_words", ())),
    )


def _save_page(path: Path, page: PageResult, meta: dict) -> None:
    """Сохраняет результат страницы атомарно (через временный файл и ``os.replace``).

    Args:
        path: Файл страницы.
        page: Результат страницы.
        meta: Параметры распознавания; страница берётся из файла, только если
            они совпадают.
    """
    data = {
        "format": PAGE_CACHE_FORMAT,
        "meta": meta,
        "page_number": page.page_number,
        "ocr_applied": page.ocr_applied,
        "plain_text": page.plain_text,
        "page_height": page.page_height,
        "review_issues": page.review_issues,
        "blocks": [_block_to_dict(b) for b in page.blocks],
    }
    part = path.with_name(path.name + ".part")
    part.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    os.replace(part, path)


def _load_page(path: Path, meta: dict) -> PageResult | None:
    """Загружает сохранённую страницу.

    Returns:
        Страница или None, если файла нет, он повреждён или получен с другими
        параметрами распознавания.
    """
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    if data.get("format") != PAGE_CACHE_FORMAT or data.get("meta") != meta:
        return None
    return PageResult(
        page_number=data["page_number"],
        ocr_applied=data["ocr_applied"],
        blocks=[_block_from_dict(b) for b in data["blocks"]],
        plain_text=data["plain_text"],
        page_height=data["page_height"],
        review_issues=data.get("review_issues", []),
    )


def _stderr_progress(done: int, total: int, page_number: int, status: str) -> None:
    """Печатает ход распознавания в stderr: «[ x%] страница N из M: статус»."""
    percent = 100 * done / total if total else 100.0
    print(
        f"[{percent:5.1f}%] страница {page_number} из {total}: {status}",
        file=sys.stderr,
        flush=True,
    )


# --- Тайминги шагов ---


def _fmt_duration(seconds: float) -> str:
    """Длительность в виде «N мин SS с» или «N с»."""
    minutes, sec = divmod(round(seconds), 60)
    return f"{minutes} мин {sec:02d} с" if minutes else f"{sec} с"


class StepTimer:
    """Учёт длительности шагов распознавания.

    Шаг относится к странице или ко всему документу (``page=None``).
    """

    def __init__(self) -> None:
        self.started = time.perf_counter()
        self.records: list[dict] = []

    @contextmanager
    def step(self, page: int | None, name: str):
        """Контекст, длительность которого записывается как шаг ``name``."""
        t0 = time.perf_counter()
        try:
            yield
        finally:
            self.add(page, name, time.perf_counter() - t0)

    def add(self, page: int | None, name: str, seconds: float) -> None:
        """Записывает шаг и печатает его длительность в stderr."""
        self.records.append({"page": page, "step": name, "seconds": round(seconds, 2)})
        where = "документ" if page is None else f"стр. {page}"
        print(f"  [{where}] {name}: {_fmt_duration(seconds)}", file=sys.stderr, flush=True)

    def report(self) -> dict:
        """Отчёт: общее время, шаги документа и шаги по страницам."""
        page_totals: dict[int, float] = {}
        for r in self.records:
            if r["page"] is not None:
                page_totals[r["page"]] = page_totals.get(r["page"], 0.0) + r["seconds"]
        return {
            "total_seconds": round(time.perf_counter() - self.started, 2),
            "document_steps": [r for r in self.records if r["page"] is None],
            "pages": {
                str(n): {
                    "total_seconds": round(page_totals[n], 2),
                    "steps": [r for r in self.records if r["page"] == n],
                }
                for n in sorted(page_totals)
            },
        }


def _print_timing_report(report: dict) -> None:
    """Печатает отчёт :meth:`StepTimer.report` в stderr."""
    lines = [f"ИТОГО по документу: {_fmt_duration(report['total_seconds'])}"]
    for step in report["document_steps"]:
        lines.append(f"  {step['step']}: {_fmt_duration(step['seconds'])}")
    for number, page in report["pages"].items():
        lines.append(f"Страница {number} — {_fmt_duration(page['total_seconds'])}")
        for step in page["steps"]:
            lines.append(f"  {step['step']}: {_fmt_duration(step['seconds'])}")
    print("\n".join(lines), file=sys.stderr, flush=True)


def _timed_fn(timer: StepTimer, page: int, name: str, fn: Callable) -> Callable:
    """Оборачивает ``fn`` так, что каждый вызов записывается как шаг ``name``."""

    def wrapped(*args, **kwargs):
        """Вызов ``fn`` с записью длительности."""
        with timer.step(page, name):
            return fn(*args, **kwargs)

    return wrapped


def ocr_document(
    pdf_path: Path | str,
    out_docx_path: Path | str,
    *,
    dpi: int = DEFAULT_DPI,
    confidence_threshold: int = DEFAULT_CONFIDENCE_THRESHOLD,
    min_page_text_chars: int = DEFAULT_MIN_PAGE_TEXT_CHARS,
    predict_tables_fn: Callable[[Path], dict] | None = None,
    image_data_fn: Callable[[Path], dict] | None = None,
    authority_name: str | None = None,
    min_free_memory_mb: int = DEFAULT_MIN_FREE_MEMORY_MB,
    table_engine: str = "tatr",
    pages_dir: Path | str | None = None,
    resume: bool = True,
    progress_fn: Callable[[int, int, int, str], None] | None = None,
    use_ai_escalation: bool = False,
    auto_rotate: bool = True,
    orientation_fn: Callable[[Path], int] | None = None,
    vlm_chat_fn: Callable[[Image.Image], str] | None = None,
    vlm_ask_fn: "vlm_reread.AskFn | None" = None,
    marks_ask_fn: Callable[[Image.Image], str] | None = None,
    diagnostic: bool = False,
    requisites_ask_fn: Callable[[Image.Image, str], str] | None = None,
) -> OcrDocumentResult:
    """Распознаёт PDF и собирает DOCX.

    Страница с текстовым слоем (не меньше ``min_page_text_chars`` знаков)
    берётся как есть. Остальные страницы рендерятся в изображения и
    распознаются (см. описание модуля). Результат каждой страницы сохраняется
    в ``pages_dir``, и при ``resume=True`` уже распознанные с теми же
    параметрами страницы не пересчитываются: прогон можно продолжить после
    остановки. Длительность шагов пишется в ``<out>.timings.json`` и в stderr.

    Движки по умолчанию создаются при первой странице-скане: документ без
    сканов не платит за загрузку моделей.

    Args:
        pdf_path: Входной PDF.
        out_docx_path: Выходной DOCX.
        dpi: Разрешение рендера страниц.
        confidence_threshold: Порог уверенности Tesseract (0–100).
        min_page_text_chars: Наименьшая длина текстового слоя страницы.
        predict_tables_fn: Таблицы через PaddleOCR (вместо движка
            ``table_engine``).
        image_data_fn: Распознавание слов Tesseract; по умолчанию
            :func:`_default_image_data_fn`.
        authority_name: Наименование органа из реквизитов документа; по нему
            восстанавливаются формула должности подписанта и шапка.
        min_free_memory_mb: Минимум свободной памяти перед каждой
            страницей-сканом, МБ; 0 — не проверять.
        table_engine: ``"tatr"`` — TATR и OCR ячеек на CPU; ``"vlm"`` —
            DeepSeek-OCR в vLLM. Если vLLM не ответил, таблицы до конца
            документа распознаются через TATR. С ``"vlm"`` таблицы,
            продолжающиеся на следующей странице, сшиваются.
        pages_dir: Каталог результатов страниц; по умолчанию ``<out>.pages``.
        resume: Брать сохранённые страницы.
        progress_fn: Ход распознавания ``(done, total, page_number, status)``;
            по умолчанию печатает в stderr.
        use_ai_escalation: Обрабатывать печати, штампы и подписи через VLM в
            Ollama. С ``table_engine="vlm"`` — удаление текста меток
            (:mod:`npa_ocr.vlm_marks`), иначе — восстановление текста под
            ними (:mod:`npa_ocr.ai_escalation`). Если Ollama недоступна,
            документ дораспознаётся без этого.
        auto_rotate: Поворачивать страницы-сканы по Tesseract OSD.
        orientation_fn: Определение поворота; по умолчанию
            :func:`_default_orientation_fn`.
        vlm_chat_fn: Запрос к DeepSeek-OCR по странице.
        vlm_ask_fn: Запрос к DeepSeek-OCR для повторного чтения ячеек.
        marks_ask_fn: Запрос к qwen3-vl для поиска меток.

    Returns:
        Страницы, путь к DOCX и тайминги.

    Raises:
        InsufficientMemoryError: Не хватило свободной памяти.
    """
    import fitz

    timer = StepTimer()
    report = progress_fn or _stderr_progress

    with timer.step(None, "Открытие и подготовка"):
        pdf_path = Path(pdf_path)
        out_docx_path = Path(out_docx_path)
        if pages_dir is None:
            pages_dir = out_docx_path.with_name(out_docx_path.stem + ".pages")
        pages_dir = Path(pages_dir)
        meta = {
            "native_layout_version": 2,
            "dpi": dpi,
            "confidence_threshold": confidence_threshold,
            "min_page_text_chars": min_page_text_chars,
            "authority_name": authority_name,
            "table_engine": table_engine,
            "auto_rotate": auto_rotate,
            "pdf_size": pdf_path.stat().st_size,
        }
        if table_engine == "vlm":
            # версия пути vlm: при её смене сохранённые страницы распознаются заново
            meta["vlm_version"] = 12
            if use_ai_escalation:
                # удаление меток (vlm_marks) вместо ИИ-эскалации
                meta["marks_version"] = 1
        if orientation_fn is None:
            orientation_fn = _default_orientation_fn
        reader = pypdf.PdfReader(str(pdf_path))
        fitz_doc = fitz.open(str(pdf_path))
        total = len(reader.pages)
        pages_dir.mkdir(parents=True, exist_ok=True)
        pages: list[PageResult] = []

    text_predict_fn: Callable | None = None
    cell_text_fn: Callable | None = None
    vlm_down: str | None = None  # table_engine="vlm": причина, по которой vLLM недоступен
    tatr_fallback_pages: set[int] = set()  # table_engine="vlm": таблицы страницы ушли в TATR
    vlm_chat_fn = vlm_chat_fn or vlm_tables.deepseek_chat
    vlm_ask_fn = vlm_ask_fn or vlm_reread.deepseek_ask
    marks_ask_fn = marks_ask_fn or vlm_marks.qwen_locate
    marks_down: str | None = None  # удаление меток: причина, по которой Ollama недоступна
    page_ds: dict[int, str] = {}  # ответ DeepSeek на страницу — для меток и второго прочтения
    source_inventory = {}
    with tempfile.TemporaryDirectory() as tmp_dir:
        for i, (pdf_page, fitz_page) in enumerate(zip(reader.pages, fitz_doc)):
            page_number = i + 1
            generation_issues = []
            from npa_ocr.source_layout import inspect_source
            source_inventory[page_number] = inspect_source(fitz_page)
            page_file = pages_dir / f"page_{page_number:04d}.json"
            if resume:
                with timer.step(page_number, "проверка сохранённого результата"):
                    cached = _load_page(page_file, meta)
                if cached is not None:
                    pages.append(cached)
                    report(page_number, total, page_number, "взято из сохранённого прогона")
                    continue
            with timer.step(page_number, "проверка текстового слоя"):
                text = (pdf_page.extract_text() or "").strip()
            mixed_native = None
            if len(text) >= min_page_text_chars:
                from npa_ocr.source_layout import extract_native_blocks
                source = source_inventory[page_number]
                native_blocks = extract_native_blocks(fitz_page, source, dpi)
                if native_blocks is not None:
                    page = PageResult(page_number=page_number, ocr_applied=False,
                                      plain_text=text,
                                      blocks=native_blocks if source['tables'] else [],
                                      page_height=fitz_page.rect.height * dpi / 72)
                    with timer.step(page_number, "сохранение страницы"):
                        _save_page(page_file, page, meta)
                    pages.append(page)
                    report(page_number, total, page_number, "текстовый слой")
                    continue
                mixed_native = extract_native_blocks(fitz_page, source, dpi, allow_raster=True)
            report(i, total, page_number, "распознаётся")
            with timer.step(page_number, "проверка свободной памяти"):
                ensure_free_memory(min_free_memory_mb, f"перед страницей {page_number}")
            use_tatr = table_engine == "tatr" and predict_tables_fn is None
            use_vlm = table_engine == "vlm" and predict_tables_fn is None
            with timer.step(page_number, "загрузка моделей"):
                if use_tatr and text_predict_fn is None:
                    text_predict_fn, cell_text_fn = _load_tatr_engines()
                if predict_tables_fn is None and not use_tatr and not use_vlm:
                    if not _paddle_runtime_usable():
                        raise RuntimeError(
                            "PaddleOCR недоступен на этом CPU (нет AVX2 / SIGILL). "
                            "Используйте table_engine='tatr' (значение по умолчанию)."
                        )
                    predict_tables_fn = _default_predict_tables_fn()
                if image_data_fn is None:
                    image_data_fn = _default_image_data_fn
            image_path = Path(tmp_dir) / f"page_{page_number}.png"
            with timer.step(page_number, "рендер страницы"):
                pix = fitz_page.get_pixmap(dpi=dpi)
                pix.save(str(image_path))
            page_height = float(pix.height)
            rotation = 0
            if auto_rotate:
                with timer.step(page_number, "поворот страницы (OSD)"):
                    rotation = orientation_fn(image_path)
                    if rotation:
                        page_height = float(rotate_page_image(image_path, rotation)[1])
            timed_image_data = _timed_fn(timer, page_number, "текст: Tesseract", image_data_fn)

            def tatr_tables_fn(path: Path, _page: int = page_number) -> list[TableBlock]:
                """Таблицы страницы через TATR с учётом времени."""
                nonlocal text_predict_fn, cell_text_fn
                if text_predict_fn is None:  # запасной путь table_engine="vlm" — модели TATR по требованию
                    with timer.step(_page, "загрузка моделей"):
                        text_predict_fn, cell_text_fn = _load_tatr_engines()
                regions_seconds = [0.0]

                def timed_regions(img: Image.Image) -> list[Box]:
                    """Поиск областей таблиц с учётом времени."""
                    t0 = time.perf_counter()
                    found = table_regions_for_page(img, text_predict_fn)
                    regions_seconds[0] = time.perf_counter() - t0
                    return found

                t0 = time.perf_counter()
                result = run_tatr_tables(
                    path,
                    text_predict_fn,
                    cell_text_fn=cell_text_fn,
                    region_fn=timed_regions,
                    release_structure=text_predict_fn.release,
                    release_cells=cell_text_fn.release,
                )
                elapsed = time.perf_counter() - t0
                timer.add(_page, "таблицы: поиск областей", regions_seconds[0])
                timer.add(_page, "таблицы: структура и текст ячеек", elapsed - regions_seconds[0])
                return result

            def vlm_tables_fn(path: Path, _page: int = page_number) -> list[TableBlock]:
                """Таблицы страницы через DeepSeek-OCR; если vLLM недоступен — через TATR."""
                # после первого отказа vLLM таблицы до конца документа идут через TATR, чтобы не
                # ждать таймаут на каждой странице; причина записывается в тайминги
                nonlocal vlm_down
                if vlm_down is None:
                    try:
                        def chat(img: Image.Image) -> str:
                            """Запрос к DeepSeek-OCR; ответ сохраняется для меток и второго прочтения."""
                            raw = vlm_chat_fn(img)
                            if getattr(raw, 'finish_reason', '') == 'length':
                                generation_issues.append(dict(
                                    reason='generation_truncated',
                                    image_size=list(img.size), usage=raw.usage))
                            # Повторный запрос по кропу таблицы имеет другую систему
                            # координат: второе чтение текста использует только страницу.
                            page_ds.setdefault(_page, raw)
                            return raw

                        return run_vlm_table_blocks(path, chat, ask_fn=vlm_ask_fn, timer=timer, page=_page)
                    except vlm_tables.VlmUnavailable as exc:
                        vlm_down = str(exc)
                        print(f"{exc} — таблицы через TATR до конца документа", file=sys.stderr, flush=True)
                        timer.add(_page, f"таблицы: vLLM недоступен, запасной путь TATR ({exc})", 0.0)
                tatr_fallback_pages.add(_page)
                return tatr_tables_fn(path)

            def clean_marks_fn(lines: list[list[vlm_marks.Word]], _page: int = page_number,
                               path: Path = image_path) -> list[list[vlm_marks.Word]]:
                """Удаляет текст меток со страницы, если на ней стоит их искать."""
                # после первого отказа Ollama метки до конца документа не удаляются
                nonlocal marks_down
                if marks_down is not None:
                    return lines
                with Image.open(path) as img:
                    img.load()
                low_boxes = [w.box for line in lines for w in line if w.conf < confidence_threshold]
                clusters = mark_detection.cluster_candidate_regions(low_boxes)
                ds_raw = page_ds.get(_page)
                if not vlm_marks.should_locate(_page, total, ds_raw, img.width, img.height, len(clusters)):
                    return lines
                t0 = time.perf_counter()
                try:
                    marks = vlm_marks.locate_marks(img, ask=marks_ask_fn)
                except vlm_marks.MarksUnavailable as exc:
                    marks_down = str(exc)
                    print(f"{exc} — без чистки меток до конца документа", file=sys.stderr, flush=True)
                    timer.add(_page, f"метки: Ollama недоступна ({exc})", 0.0)
                    return lines
                finally:
                    timer.add(_page, "метки: qwen3-vl (Ollama)", time.perf_counter() - t0)
                print(f"  [стр. {_page}] метки: " + (", ".join(m.label for m in marks) or "нет"),
                      file=sys.stderr, flush=True)
                return vlm_marks.clean_lines(lines, marks, ds_raw, img.width, img.height, authority_name)

            def reread_text_fn(lines: list[list[vlm_marks.Word]], _page: int = page_number,
                               path: Path = image_path) -> tuple[list[list[vlm_marks.Word]], list[vlm_text.Formula]]:
                """Второе прочтение текста по ответу DeepSeek, полученному для таблиц."""
                with Image.open(path) as img:
                    width, height = img.size
                with timer.step(_page, "текст: второе прочтение DeepSeek"):
                    restored, formulas = vlm_text.reread_lines(lines, page_ds.get(_page), width, height)
                    from npa_ocr.requisites import restore_header_number, restore_inline_requisites
                    from npa_ocr.text_recovery import recover_titles, recover_inline_codes

                    with Image.open(path) as img:
                        restored = recover_titles(restored, page_ds.get(_page), img)
                        restored = restore_header_number(restored, page_ds.get(_page), img, vlm_ask_fn)
                        restored = restore_inline_requisites(restored, img, vlm_ask_fn,
                            requisites_ask_fn or (vlm_marks.qwen_read if use_ai_escalation else None))
                        restored = recover_inline_codes(restored, img, vlm_ask_fn)
                    return restored, formulas

            use_marks = use_vlm and use_ai_escalation
            if use_tatr or use_vlm:
                blocks = ocr_page(
                    image_path,
                    table_blocks_fn=vlm_tables_fn if use_vlm else tatr_tables_fn,
                    image_data_fn=timed_image_data,
                    confidence_threshold=confidence_threshold,
                    authority_name=authority_name,
                    detect_marks_fn=_default_detect_marks_fn if use_ai_escalation and not use_marks else None,
                    recover_covered_text_fn=(
                        ai_escalation.recover_covered_text if use_ai_escalation and not use_marks else None),
                    clean_marks_fn=clean_marks_fn if use_marks else None,
                    reread_text_fn=reread_text_fn if use_vlm else None,
                )
            else:
                blocks = ocr_page(
                    image_path,
                    predict_tables_fn=_timed_fn(timer, page_number, "таблицы: PaddleOCR", predict_tables_fn),
                    image_data_fn=timed_image_data,
                    confidence_threshold=confidence_threshold,
                    authority_name=authority_name,
                    detect_marks_fn=_default_detect_marks_fn if use_ai_escalation else None,
                    recover_covered_text_fn=ai_escalation.recover_covered_text if use_ai_escalation else None,
                )
            if mixed_native is not None and not rotation:
                from npa_ocr.source_layout import merge_mixed_blocks
                blocks = merge_mixed_blocks(mixed_native, blocks,
                                            source_inventory[page_number], dpi)
            page = PageResult(
                page_number=page_number,
                ocr_applied=True,
                blocks=blocks,
                page_height=page_height,
                review_issues=generation_issues,
            )
            # страница с таблицами из запасного TATR сохраняется с другими параметрами, чтобы
            # следующий прогон распознал её через vLLM заново
            page_meta = {**meta, "tables": "tatr"} if page_number in tatr_fallback_pages else meta
            with timer.step(page_number, "сохранение страницы"):
                _save_page(page_file, page, page_meta)
            pages.append(page)
            report(page_number, total, page_number, "готово")

    with timer.step(None, "Завершение распознавания"):
        fitz_doc.close()
    with timer.step(None, "Колонтитулы"):
        pages = remove_running_headers_and_footers(pages)
    if table_engine == "vlm":
        from npa_ocr.coverage import remove_table_duplicates
        pages = [remove_table_duplicates(page) for page in pages]
        # пороги сшивки подобраны под рамки таблиц vlm_tables; у TATR рамки другие
        with timer.step(None, "Сшивка таблиц между страницами"):
            pages, stitch_log = stitch_tables(pages)
        for line in stitch_log:
            print(f"  [сшивка] {line}", file=sys.stderr, flush=True)
    with timer.step(None, "Сборка DOCX"):
        build_docx(pages, out_docx_path, dpi=dpi, diagnostic=diagnostic)
        from npa_ocr.coverage import review_report
        out_docx_path.with_suffix(".review.json").write_text(
            json.dumps(review_report(pages, source_inventory=source_inventory),
                       ensure_ascii=False, indent=2), encoding="utf-8")

    timings = timer.report()
    timings_path = out_docx_path.with_name(out_docx_path.stem + ".timings.json")
    timings_path.write_text(json.dumps(timings, ensure_ascii=False, indent=2), encoding="utf-8")
    _print_timing_report(timings)
    return OcrDocumentResult(pages=pages, docx_path=out_docx_path, timings=timings, timings_path=timings_path)
